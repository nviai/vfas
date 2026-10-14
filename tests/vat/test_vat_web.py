"""Tax web.

Invoice lists, returns (save, finalize, reopen, XML, PDF, Excel), VAT
offset, CIT, calendar.
"""

import io

from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.registry import ACCOUNTANT, CASHIER, VIEWER
from tests.conftest import ROLE_EMAILS
from tests.vat.test_vat import parties, q1  # noqa: F401  (fixture)


def test_vat_pages_and_exports(web, users, q1):  # noqa: F811
    web.login(ROLE_EMAILS[ACCOUNTANT])
    base = "/vat/declarations/quarter/2026/1"
    r = web.get(base)
    assert (
        r.status_code == 200 and "860.000" in r.text and "13.000.000" in r.text
    )
    for url in (
        "/vat/declarations",
        "/vat/lists?direction=out&period_type=quarter&year=2026&period=1",
        "/vat/lists?direction=in&period_type=quarter&year=2026&period=1",
        "/vat/cit?year=2026",
        "/vat/calendar",
        "/vat/lists",
    ):
        assert web.get(url).status_code == 200, url
    xml = web.get(base + ".xml")
    assert xml.status_code == 200 and b"<ct40>860000</ct40>" in xml.content
    pdf = web.get(base + ".pdf")
    assert pdf.status_code == 200
    text = "".join(
        p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages
    )
    assert "TỜ KHAI THUẾ GIÁ TRỊ GIA TĂNG" in text and "860.000" in text
    assert web.get(base + "/reduction.pdf").status_code == 200
    x = web.get(base + ".xlsx")
    assert (
        x.status_code == 200
        and load_workbook(io.BytesIO(x.content)).active.max_row > 20
    )
    for d in ("in", "out"):
        q = f"?direction={d}&period_type=quarter&year=2026&period=1"
        assert web.get("/vat/lists.xlsx" + q).status_code == 200
        assert web.get("/vat/lists.pdf" + q).status_code == 200
    assert web.get("/vat/cit.pdf?year=2026").status_code == 200
    assert web.get("/vat/cit.xlsx?year=2026").status_code == 200
    # Save with manual adjustments, finalize, reopen
    r = web.post(
        base,
        {
            "ct37": "100.000",
            "ct38": "0",
            "ct39a": "",
            "ct40b": "",
            "ct42": "",
            "action": "save",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    assert "960.000" in web.get(base).text
    r = web.post(
        base, {"ct37": "100000", "action": "final"}, follow_redirects=False
    )
    assert r.status_code == 303 and "Đã chốt" in web.get(base).text
    assert (
        web.post(base + "/reopen", follow_redirects=False).status_code == 303
    )
    # VAT offset entry
    r = web.post(base + "/offset", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith(
        "/entries/"
    )
    # CIT
    r = web.post(
        "/vat/cit",
        {"year": "2026", "B1": "1.000.000", "rate": "20", "action": "save"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    web.get("/lang/en")
    assert (
        web.get(base).status_code == 200
        and web.get("/vat/cit?year=2026").status_code == 200
    )


def test_vat_permissions(web, users):
    web.login(ROLE_EMAILS[VIEWER])
    assert web.get("/vat/declarations/quarter/2026/1").status_code == 200
    web.get("/vat/declarations/quarter/2026/1")
    assert (
        web.post(
            "/vat/declarations/quarter/2026/1", {"action": "save"}
        ).status_code
        == 403
    )
    web.login(ROLE_EMAILS[CASHIER])
    assert web.get("/vat/lists").status_code == 403
