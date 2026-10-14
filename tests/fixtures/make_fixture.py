"""Generate a sample cash book Excel file for the import tests.

The file mirrors a cash book used in real life. Run::

    .venv/bin/python -m tests.fixtures.make_fixture

to write tests/fixtures/so_quy_mau.xlsx.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent
FIXTURE = HERE / "so_quy_mau.xlsx"

COMPANY = "CÔNG TY TNHH CÔNG NGHỆ THÔNG MINH NVIAI"
ADDRESS = (
    "Tầng 14, Tòa nhà HM Town, 412 Nguyễn Thị Minh Khai, Phường 05, "
    "Quận 3, TP. Hồ Chí Minh"
)

OPENING = 5_000_000

# (date, kind, voucher number, amount, description, contra account)
ROWS: list[tuple[date, str, str, int, str, str]] = [
    (
        date(2023, 1, 12),
        "thu",
        "PT0001",
        50_000_000,
        "Rút tiền gửi ngân hàng về nhập quỹ",
        "1121",
    ),
    (
        date(2023, 1, 15),
        "chi",
        "PC0001",
        3_500_000,
        "Mua văn phòng phẩm",
        "6422",
    ),
    (
        date(2023, 2, 1),
        "chi",
        "PC0002",
        12_000_000,
        "Trả lương tháng 01/2023",
        "334",
    ),
    (
        date(2023, 2, 10),
        "thu",
        "PT0002",
        8_800_000,
        "Thu tiền bán hàng HĐ 0000123",
        "5111",
    ),
    (
        date(2023, 3, 5),
        "chi",
        "PC0003",
        1_100_000,
        "Thuế GTGT đầu vào dịch vụ sửa chữa",
        "1331",
    ),
    (
        date(2023, 4, 12),
        "chi",
        "PC0004",
        5_000_000,
        "Tạm ứng công tác phí",
        "141",
    ),
    (
        date(2023, 5, 20),
        "thu",
        "PT0003",
        15_000_000,
        "Thu nợ khách hàng Công ty ABC",
        "131",
    ),
    (
        date(2023, 6, 30),
        "chi",
        "PC0005",
        2_350_000,
        "Chi phí tiếp khách",
        "6422",
    ),
    (
        date(2023, 8, 15),
        "chi",
        "PC0006",
        9_000_000,
        "Trả tiền hàng cho NCC Minh Phát",
        "331",
    ),
    (
        date(2023, 10, 2),
        "thu",
        "PT0004",
        20_000_000,
        "Nhận vốn góp bổ sung",
        "4111",
    ),
    (
        date(2023, 12, 28),
        "chi",
        "PC0007",
        4_200_000,
        "Nộp bảo hiểm xã hội",
        "3383",
    ),
    (
        date(2024, 1, 10),
        "thu",
        "PT0005",
        30_000_000,
        "Rút tiền gửi ngân hàng về nhập quỹ",
        "1121",
    ),
    (
        date(2024, 2, 5),
        "chi",
        "PC0008",
        18_500_000,
        "Trả lương tháng 01/2024",
        "334",
    ),
    (
        date(2024, 3, 18),
        "chi",
        "PC0009",
        750_000,
        "Phí chuyển phát nhanh",
        "6422",
    ),
    (
        date(2024, 5, 9),
        "thu",
        "PT0006",
        6_600_000,
        "Thu tiền cung cấp dịch vụ",
        "5113",
    ),
    (
        date(2024, 7, 1),
        "chi",
        "PC0010",
        3_000_000,
        "Nộp lệ phí môn bài",
        "3338",
    ),
    (date(2024, 9, 14), "chi", "PC0011", 1_250_000, "Trả lãi vay", "635"),
    (
        date(2024, 11, 25),
        "thu",
        "PT0007",
        2_000_000,
        "Hoàn ứng tiền tạm ứng thừa",
        "141",
    ),
    (
        date(2024, 12, 31),
        "chi",
        "PC0012",
        11_000_000,
        "Trả lương tháng 12/2024",
        "334",
    ),
    (
        date(2025, 1, 15),
        "thu",
        "PT0008",
        25_000_000,
        "Rút tiền gửi ngân hàng về nhập quỹ",
        "1121",
    ),
    (
        date(2025, 3, 3),
        "chi",
        "PC0013",
        4_400_000,
        "Mua công cụ dụng cụ",
        "153",
    ),
    (date(2025, 4, 22), "chi", "PC0014", 600_000, "Chi phí khác", "811"),
    (
        date(2025, 6, 10),
        "thu",
        "PT0009",
        1_500_000,
        "Thu thanh lý phế liệu",
        "711",
    ),
    (
        date(2025, 8, 8),
        "chi",
        "PC0015",
        7_700_000,
        "Trả nợ vay ngắn hạn",
        "341",
    ),
    (
        date(2025, 9, 30),
        "chi",
        "PC0016",
        2_000_000,
        "Nộp thuế thu nhập cá nhân",
        "3335",
    ),
]

EXAMPLE = (
    date(2023, 1, 12),
    "thu",
    "PT0000",
    1_000_000,
    "Thu tiền bán hàng (ví dụ cách ghi)",
    "5111",
)


def _header(
    ws, title: str, form: str, columns: list[str], letters: list[str]
) -> None:
    ws["A1"] = f"Đơn vị: {COMPANY}"
    ws["A2"] = f"Địa chỉ: {ADDRESS}"
    ws.cell(1, len(columns) - 2, f"Mẫu số {form}").font = Font(bold=True)
    ws.cell(
        2,
        len(columns) - 2,
        "(Ban hành theo Thông tư số 133/2016/TT-BTC ngày 26/8/2016 "
        "của Bộ Tài chính)",
    )
    ws["A4"] = title
    ws["A4"].font = Font(bold=True, size=14)
    ws["A5"] = "Loại quỹ: Tiền Việt Nam"
    ws["A6"] = "Từ ngày 12/01/2023 đến ngày 30/09/2025"
    ws["A7"] = "Đơn vị tính: đồng"
    for col, (name, letter) in enumerate(zip(columns, letters), start=1):
        ws.cell(9, col, name).alignment = Alignment(
            wrap_text=True, horizontal="center"
        )
        ws.cell(11, col, letter).alignment = Alignment(horizontal="center")


def write_book(
    path: Path,
    rows,
    opening=OPENING,
    *,
    example=None,
    contra: bool = True,
    totals: bool = True,
    string_dates_every: int = 3,
    opening_cell=None,
) -> Path:
    """Write a cash book file.

    rows: [(post_date, kind 'thu'/'chi', number, amount, description,
    contra)]. Elements may be None or "bad" values to test errors. The
    optional 7th element is a dict of cell overrides: {"doc_date": ...,
    "receipt": ..., "payment": ..., "receipt_no": ..., "payment_no":
    ..., "note": ...}.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "S04a-DNN"
    _header(
        ws,
        "SỔ QUỸ TIỀN MẶT",
        "S04a-DNN",
        [
            "Ngày, tháng ghi sổ",
            "Ngày, tháng chứng từ",
            "Số hiệu chứng từ Thu",
            "Số hiệu chứng từ Chi",
            "Diễn giải",
            "Thu",
            "Chi",
            "Tồn",
            "Ghi chú",
        ],
        ["A", "B", "C", "D", "E", "1", "2", "3", "G"],
    )
    ws["E12"] = "Số tồn đầu kỳ"
    ws["H12"] = opening if opening_cell is None else opening_cell

    wb_b = wb.create_sheet("S04b-DNN") if contra else None
    if wb_b is not None:
        _header(
            wb_b,
            "SỔ KẾ TOÁN CHI TIẾT QUỸ TIỀN MẶT",
            "S04b-DNN",
            [
                "Ngày, tháng ghi sổ",
                "Ngày, tháng chứng từ",
                "Số hiệu chứng từ Thu",
                "Số hiệu chứng từ Chi",
                "Diễn giải",
                "TK đối ứng",
                "Nợ",
                "Có",
                "Tồn",
                "Ghi chú",
            ],
            ["A", "B", "C", "D", "E", "F", "1", "2", "3", "G"],
        )
        wb_b["E12"] = "Số tồn đầu kỳ"
        wb_b["I12"] = opening if isinstance(opening, int) else None

    r = 13
    bal = opening if isinstance(opening, int) else 0
    total_in = total_out = 0
    all_rows = (
        [(*example, {"note": "Dòng ví dụ"})] if example else []
    ) + list(rows)
    for i, item in enumerate(all_rows):
        d, kind, number, amount, desc, acc = item[:6]
        over = item[6] if len(item) > 6 else {}
        is_example = "Dòng ví dụ" in (over.get("note") or "")
        as_string = (
            string_dates_every
            and i % string_dates_every == 0
            and isinstance(d, date)
        )
        post_cell = (
            d.strftime("%d/%m/%Y")
            if as_string
            else (
                datetime(d.year, d.month, d.day) if isinstance(d, date) else d
            )
        )
        doc = over.get("doc_date", post_cell)
        rec = over.get("receipt", amount if kind == "thu" else None)
        pay = over.get("payment", amount if kind == "chi" else None)
        rec_no = over.get("receipt_no", number if kind == "thu" else None)
        pay_no = over.get("payment_no", number if kind == "chi" else None)
        if not is_example and isinstance(rec, int):
            bal += rec
            total_in += rec
        if not is_example and isinstance(pay, int):
            bal -= pay
            total_out += pay
        values = [
            post_cell,
            doc,
            rec_no,
            pay_no,
            desc,
            rec,
            pay,
            over.get("balance", bal),
            over.get("note"),
        ]
        for col, v in enumerate(values, start=1):
            ws.cell(r, col, v)
        if wb_b is not None:
            for col, v in enumerate(
                [
                    post_cell,
                    doc,
                    rec_no,
                    pay_no,
                    desc,
                    acc,
                    rec,
                    pay,
                    bal,
                    over.get("note"),
                ],
                start=1,
            ):
                wb_b.cell(r, col, v)
        r += 1
    if totals:
        ws.cell(r, 5, "Cộng phát sinh trong kỳ")
        ws.cell(r, 6, total_in)
        ws.cell(r, 7, total_out)
        ws.cell(r + 1, 5, "Số tồn cuối kỳ")
        ws.cell(r + 1, 8, bal)
        ws.cell(
            r + 3, 1, "Sổ này có 02 trang, đánh số từ trang 01 đến trang 02"
        )
        ws.cell(r + 4, 1, "Ngày mở sổ: 12/01/2023")
        if wb_b is not None:
            wb_b.cell(r, 5, "Cộng phát sinh trong kỳ")
            wb_b.cell(r, 7, total_in)
            wb_b.cell(r, 8, total_out)
            wb_b.cell(r + 1, 5, "Số tồn cuối kỳ")
            wb_b.cell(r + 1, 9, bal)
    for sheet in [ws] + ([wb_b] if wb_b is not None else []):
        for row in sheet.iter_rows(min_row=12):
            for c in row:
                if isinstance(c.value, int):
                    c.number_format = "#,##0"
                elif isinstance(c.value, datetime):
                    c.number_format = "dd/mm/yyyy"
        sheet.column_dimensions["E"].width = 40
    wb.save(path)
    return path


def expected() -> dict:
    total_in = sum(a for _, k, _, a, _, _ in ROWS if k == "thu")
    total_out = sum(a for _, k, _, a, _, _ in ROWS if k == "chi")
    return {
        "rows": len(ROWS),
        "total_in": total_in,
        "total_out": total_out,
        "closing": OPENING + total_in - total_out,
        "opening": OPENING,
    }


def main() -> None:
    write_book(FIXTURE, ROWS, OPENING, example=EXAMPLE)
    print(f"Đã ghi {FIXTURE} – {expected()}")


if __name__ == "__main__":
    main()
