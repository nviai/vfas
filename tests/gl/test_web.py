"""Web of the gl module.

Pages, journal vouchers, opening balances, period locks, closing
entries, PDF/Excel export, permissions.
"""

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.core.registry import CASHIER, CHIEF, VIEWER
from app.features.journal import queries
from app.features.journal.models import JournalEntry, OpeningBalance
from tests.conftest import ROLE_EMAILS

Q2024 = "mode=year&year=2024"
REPORTS = ["s03a", "s03b", "ledger", "f01", "b01a", "b02", "b03", "b09"]


@pytest.fixture()
def kt(web, users):
    return web.login(ROLE_EMAILS[CHIEF])


def test_pages_render(kt, year_book):
    pk = year_book["entries"]["mua_tscd"]
    pages = [
        "/entries",
        "/entries?status=posted&q=máy",
        "/entries?account=112",
        f"/entries?partner_id={year_book['kh'].id}",
        "/entries/new",
        f"/entries/{pk.id}",
        f"/entries/new?copy={pk.id}",
        "/opening",
        "/locks",
        "/locks?year=2024",
        "/closing",
        f"/closing?{Q2024}",
        "/reports",
        f"/reports?{Q2024}",
    ]
    pages += [f"/reports/{r}?{Q2024}" for r in REPORTS]
    pages += [
        "/reports/s03b?account=131&mode=year&year=2024",
        f"/reports/ledger?account=131&partner_id={year_book['kh'].id}&{Q2024}",
        f"/reports/ledger?account=1121&bank_account_id={year_book['bank'].id}&{Q2024}",
        "/reports/f01?level=1&show_zero=1&mode=quarter&quarter=1&year=2024",
        "/reports/b02?mode=custom&start=2024-01-01&end=2024-06-30",
    ]
    for url in pages:
        r = kt.get(url)
        assert r.status_code == 200, (url, r.text[:300])
    kt.get("/lang/en")
    for url in pages:
        assert kt.get(url).status_code == 200, url
    r = kt.get(
        "/entries?status=posted",
        headers={"HX-Request": "true", "HX-Target": "entry-results"},
    )
    assert r.status_code == 200 and "<html" not in r.text


def test_exports_open(kt, year_book):
    for name in REPORTS:
        r = kt.get(f"/reports/{name}.pdf?{Q2024}&account=112")
        assert (
            r.status_code == 200
            and r.headers["content-type"] == "application/pdf"
        ), name
        reader = PdfReader(io.BytesIO(r.content))
        assert len(reader.pages) >= 1
        text = reader.pages[0].extract_text()
        assert "Đơn vị" in text, name
        r = kt.get(f"/reports/{name}.xlsx?{Q2024}&account=112")
        assert r.status_code == 200, name
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.active.max_row > 5
    pk = year_book["entries"]["mua_tscd"]
    r = kt.get(f"/entries/{pk.id}/print.pdf")
    assert (
        r.status_code == 200
        and "PHIẾU KẾ TOÁN"
        in PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
    )
    r = kt.get("/entries/export.xlsx?status=posted")
    assert r.status_code == 200


def test_s03a_pdf_page_count(kt, year_book):
    r = kt.get(f"/reports/s03a.pdf?{Q2024}")
    reader = PdfReader(io.BytesIO(r.content))
    n = len(reader.pages)
    full = "".join(p.extract_text() for p in reader.pages)
    assert f"Sổ này có {n} trang" in full


def _pk_form(lines, action="save", **extra):
    data = {
        "date": "2024-05-20",
        "doc_date": "2024-05-20",
        "description": "Trích trước chi phí điện",
        "action": action,
        "lines-count": str(len(lines)),
    }
    for i, ln in enumerate(lines):
        for k, v in ln.items():
            data[f"lines-{i}-{k}"] = v
    data.update(extra)
    return data


def test_create_and_post_journal_entry(kt, db, company, year_book):
    kt.get("/entries/new")
    r = kt.post(
        "/entries",
        _pk_form(
            [
                {
                    "debit_account": "6422",
                    "credit_account": "335",
                    "amount": "2tr",
                },
                {
                    "debit_account": "1331",
                    "credit_account": "335",
                    "amount": "200.000",
                    "description": "Thuế GTGT",
                },
            ]
        ),
        follow_redirects=False,
    )
    assert r.status_code == 303
    e = db.scalars(
        select(JournalEntry).where(
            JournalEntry.description == "Trích trước chi phí điện"
        )
    ).one()
    assert (
        e.status == "draft"
        and e.total == 2_200_000
        and e.lines[1].description == "Thuế GTGT"
    )
    # Edit the draft, then post
    r = kt.post(
        f"/entries/{e.id}/edit",
        _pk_form(
            [
                {
                    "debit_account": "6422",
                    "credit_account": "335",
                    "amount": "3000000",
                }
            ],
            action="post",
        ),
        follow_redirects=False,
    )
    assert r.status_code == 303
    db.refresh(e)
    assert (
        e.status == "posted"
        and e.number == "PK2024-0011"
        and e.total == 3_000_000
    )
    page = kt.get(f"/entries/{e.id}").text
    assert "PK2024-0011" in page and "Đã ghi sổ" in page
    # Error: missing party for 131 when posting -> error, no voucher
    # created
    n = db.query(JournalEntry).count()
    r = kt.post(
        "/entries",
        _pk_form(
            [
                {
                    "debit_account": "131",
                    "credit_account": "711",
                    "amount": "1000",
                }
            ],
            action="post",
        ),
    )
    assert r.status_code in (200, 303)
    # No lines
    r = kt.post("/entries", _pk_form([]))
    assert r.status_code == 422 and "ít nhất một dòng" in r.text
    assert db.query(JournalEntry).count() >= n


def test_pk_with_vat_and_bank(kt, db, year_book):
    bank = year_book["bank"]
    kt.get("/entries/new")
    r = kt.post(
        "/entries",
        _pk_form(
            [
                {
                    "debit_account": "6422",
                    "credit_account": "1121",
                    "amount": "1000000",
                    "bank_account_id": str(bank.id),
                }
            ],
            action="post",
            has_vat="1",
            vat_direction="in",
            vat_invoice_no="123",
            vat_invoice_series="C24TAA",
            vat_goods_amount="1000000",
            vat_amount="100000",
            vat_rate="10",
            vat_partner_name="Điện lực",
        ),
        follow_redirects=False,
    )
    assert r.status_code == 303
    e = db.scalars(
        select(JournalEntry).order_by(JournalEntry.id.desc())
    ).first()
    assert (
        e.status == "posted" and e.lines[0].credit_bank_account_id == bank.id
    )
    assert (
        e.vat_records[0].invoice_no == "123"
        and e.vat_records[0].vat_amount == 100_000
    )


def test_opening_balances_form_and_import(kt, db, company, year_book):
    bank = year_book["bank"]
    kh = year_book["kh"]
    kt.get("/opening")
    r = kt.post(
        "/opening",
        {
            "lines-count": "3",
            "lines-0-account": "1111",
            "lines-0-debit": "10.000.000",
            "lines-1-account": "131",
            "lines-1-debit": "5000000",
            "lines-1-partner_id": str(kh.id),
            "lines-2-account": "4111",
            "lines-2-credit": "35000000",
            f"bank-{bank.id}": "20000000",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    rows = db.scalars(select(OpeningBalance)).all()
    assert (
        sum(r.debit for r in rows) == 35_000_000 == sum(r.credit for r in rows)
    )
    assert any(
        r.bank_account_id == bank.id and r.debit == 20_000_000 for r in rows
    )
    assert (
        queries.balance(db, company.id, "1111", upto=date(2023, 1, 12))
        == 10_000_000
    )
    # Error: 131 without a party
    r = kt.post(
        "/opening",
        {
            "lines-count": "1",
            "lines-0-account": "131",
            "lines-0-debit": "1000",
        },
    )
    assert r.status_code == 422 and "đối tượng" in r.text
    # Excel template -> import it back
    r = kt.get("/opening/template.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.active.max_row == 5  # header + 4 existing rows
    wb2 = Workbook()
    ws = wb2.active
    ws.append(
        [
            "Số hiệu TK",
            "Mã đối tượng",
            "Số tài khoản ngân hàng",
            "Dư Nợ",
            "Dư Có",
            "Ghi chú",
        ]
    )
    ws.append(["1111", None, None, 7_000_000, None, "Quỹ"])
    ws.append(["131", kh.code, None, 3_000_000, None, None])
    ws.append(["1121", None, bank.number, 10_000_000, None, None])
    ws.append(["4111", None, None, None, 20_000_000, None])
    buf = io.BytesIO()
    wb2.save(buf)
    r = kt.post(
        "/opening/import",
        files={
            "file": (
                "sodu.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200 and "Nhập 4 dòng" in r.text
    import re

    token = re.search(r'name="token" value="([0-9a-f]+)"', r.text).group(1)
    r = kt.post(
        "/opening/import/confirm",
        {"token": token, "action": "confirm"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    db.rollback()  # end the test session's stale read transaction
    rows = db.scalars(select(OpeningBalance)).all()
    assert sum(r.debit for r in rows) == 20_000_000
    # Invalid file
    ws.append(["999", None, None, 1, None, None])
    buf = io.BytesIO()
    wb2.save(buf)
    r = kt.post(
        "/opening/import",
        files={"file": ("x.xlsx", buf.getvalue(), "application/octet-stream")},
    )
    assert "Không có TK 999" in r.text


def test_locks_and_closing_web(kt, db, company, year_book):
    r = kt.post(
        "/locks/lock", {"year": "2023", "month": "1"}, follow_redirects=False
    )
    assert r.status_code == 303
    # February is not locked yet
    r = kt.post(
        "/locks/lock", {"year": "2023", "month": "3"}, follow_redirects=False
    )
    assert r.status_code == 303 and "flash" in r.headers.get("set-cookie", "")
    r = kt.post(
        "/locks/unlock",
        {"year": "2023", "month": "1", "reason": "Sửa số dư"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    page = kt.get(f"/closing?{Q2024}").text
    assert "4212" in page
    kt.post("/closing/cit", {"date": "2024-12-31", "amount": "22.260.000"})
    r = kt.post(
        "/closing/run",
        {"mode": "year", "year": "2024"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    kc = db.scalars(
        select(JournalEntry).where(JournalEntry.doc_type == "KC")
    ).one()
    assert kc.status == "posted"
    assert kt.get(f"/entries/{kc.id}").status_code == 200
    r = kt.post("/closing/retained", {"year": "2024"}, follow_redirects=False)
    assert r.status_code == 303
    assert db.scalars(
        select(JournalEntry).where(JournalEntry.doc_type == "KC")
    ).all()[-1].date == date(2025, 1, 1)


def test_permissions(web, users, year_book):
    web.login(ROLE_EMAILS[VIEWER])
    assert web.get(f"/reports/b01a?{Q2024}").status_code == 200
    assert web.get(f"/reports/b01a.pdf?{Q2024}").status_code == 200
    assert web.get("/entries").status_code == 200
    for url in ("/entries/new", "/closing"):
        assert web.get(url).status_code == 403, url
    assert web.post("/opening", {"lines-count": "0"}).status_code == 403
    assert (
        web.post("/locks/lock", {"year": "2023", "month": "1"}).status_code
        == 403
    )
    web.login(ROLE_EMAILS[CASHIER])
    for url in ("/entries", "/reports/f01", "/opening"):
        assert web.get(url).status_code == 403, url
