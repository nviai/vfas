"""Books, tax and payroll end to end.

Reports (web + PDF + Excel), 01/GTGT return + XML, period lock/unlock,
payroll -> posting -> salary payment.
"""

import io
import re
import xml.etree.ElementTree as ET
from datetime import date

from openpyxl import load_workbook
from playwright.sync_api import expect
from pypdf import PdfReader

from tests.e2e.conftest import pick


def _check_exports(s, expect_pdf_text: str | None = None) -> None:
    links = s.page.locator(
        "main a[href*='.pdf'], main a[href*='.xlsx']"
    ).evaluate_all("els => els.map(e => e.getAttribute('href'))")
    kinds = {".pdf": False, ".xlsx": False}
    for href in links:
        r = s.get(href)
        assert r.status == 200, href
        if ".pdf" in href:
            assert r.headers["content-type"].startswith("application/pdf"), (
                href
            )
            text = "".join(
                p.extract_text() for p in PdfReader(io.BytesIO(r.body())).pages
            )
            if expect_pdf_text:
                assert expect_pdf_text in text, href
            kinds[".pdf"] = True
        else:
            assert "spreadsheetml" in r.headers["content-type"], href
            assert load_workbook(io.BytesIO(r.body())).sheetnames
            kinds[".xlsx"] = True
    assert all(kinds.values()), f"missing export buttons: {links}"


def test_financial_reports_web_pdf_excel(owner):
    s = owner
    last_year = date.today().year - 1
    s.goto(f"/reports/b01a?mode=year&year={last_year}")
    main = s.page.locator("main")
    expect(main).to_contain_text(
        re.compile("BÁO CÁO TÌNH HÌNH TÀI CHÍNH", re.I)
    )
    expect(main).to_contain_text("Tổng tài sản (200) = Tổng nguồn vốn (500)")
    _check_exports(s, "BÁO CÁO TÌNH HÌNH TÀI CHÍNH")

    s.goto(f"/reports/f01?mode=year&year={last_year}")
    expect(s.page.locator("main")).to_contain_text(
        re.compile("BẢNG CÂN ĐỐI TÀI KHOẢN", re.I)
    )
    _check_exports(s)

    s.goto("/reports/s03b")
    # general ledger by level-1 account
    pick(
        s, s.page.locator("main input[placeholder='Gõ để tìm…']").first, "131"
    )
    s.page.locator("main button[type=submit]").first.click()
    s.page.wait_for_load_state("networkidle")
    expect(s.page.locator("main")).to_contain_text(re.compile("SỔ CÁI", re.I))
    expect(s.page.locator("main")).to_contain_text(
        "131 – Phải thu của khách hàng"
    )
    _check_exports(s)

    s.goto("/cash/book")
    expect(s.page.locator("main")).to_contain_text(
        re.compile("SỔ QUỸ TIỀN MẶT", re.I)
    )
    _check_exports(s, "SỔ QUỸ TIỀN MẶT")


def test_vat_declaration_view_and_xml(owner):
    s = owner
    s.goto("/vat/declarations")
    last_year = date.today().year - 1
    s.page.locator(
        f"main a[href='/vat/declarations/quarter/{last_year}/4']"
    ).click()
    s.page.wait_for_load_state("networkidle")
    main = s.page.locator("main")
    expect(main).to_contain_text("01/GTGT")
    expect(main).to_contain_text("[43]")
    href = s.page.locator("main a[href*='.xml']").first.get_attribute("href")
    r = s.get(href)
    assert r.status == 200 and "xml" in r.headers["content-type"]
    root = ET.fromstring(r.body())
    assert root.tag.endswith("HSoThueDTu")


def test_lock_and_unlock_with_reason(owner):
    s = owner
    year = date.today().year
    s.goto(f"/locks?year={year}")
    lock_btn = s.page.locator("main form[action='/locks/lock'] button").first
    lock_btn.click()
    confirm = s.page.get_by_role("alertdialog")
    confirm.locator("button[x-ref=ok]").click()
    s.page.wait_for_load_state("networkidle")
    expect(s.page.locator("main")).to_contain_text(
        f"Đã khóa đến hết tháng 01/{year}"
    )

    # Unlock: a reason is required (drop the browser's required
    # attribute to test the server-side check)
    s.page.get_by_role("button", name="Mở khóa").first.click()
    dialog = s.page.get_by_role("dialog").filter(has_text="Lý do mở khóa")
    dialog.locator("textarea[name=reason]").evaluate(
        "el => el.removeAttribute('required')"
    )
    dialog.get_by_role("button", name="Mở khóa").click()
    s.page.wait_for_load_state("networkidle")
    expect(s.page.locator("body")).to_contain_text("bắt buộc phải nhập lý do")
    expect(s.page.locator("main")).to_contain_text(
        f"Đã khóa đến hết tháng 01/{year}"
    )
    s.page.get_by_role("button", name="Mở khóa").first.click()
    dialog.locator("textarea[name=reason]").fill("Kiểm thử E2E mở lại kỳ")
    dialog.get_by_role("button", name="Mở khóa").click()
    s.page.wait_for_load_state("networkidle")
    expect(s.page.locator("main")).not_to_contain_text(
        f"Đã khóa đến hết tháng 01/{year}"
    )
    s.goto("/audit?group=period")
    expect(s.page.locator("main")).to_contain_text("Kiểm thử E2E mở lại kỳ")


def test_payroll_create_post_and_pay(accountant):
    s = accountant
    today = date.today()
    s.goto("/payroll")
    existing = s.page.locator(
        "main a[href^='/payroll/runs/']",
        has_text=f"{today.month:02d}/{today.year}",
    )
    if existing.count():
        existing.first.click()
    else:
        s.page.get_by_role("button", name="Lập bảng lương").first.click()
        dialog = s.page.get_by_role("dialog")
        dialog.locator("select[name=month]").select_option(str(today.month))
        dialog.locator("input[name=year]").fill(str(today.year))
        dialog.get_by_role("button", name="Lập bảng lương").click()
    s.page.wait_for_url(re.compile(r"/payroll/runs/\d+$"))
    main = s.page.locator("main")
    expect(main).to_contain_text("Thực lĩnh")
    if s.page.get_by_role("button", name=re.compile("Lưu và ghi sổ")).count():
        s.page.get_by_role("button", name=re.compile("Lưu và ghi sổ")).click()
        s.page.get_by_role("alertdialog").locator("button[x-ref=ok]").click()
        s.page.wait_for_load_state("networkidle")
    expect(main).to_contain_text("Đã ghi sổ")
    expect(main).to_contain_text(re.compile(r"LG\d{4}-\d{4}"))
    if s.page.get_by_role("button", name="Trả lương").count():
        s.page.get_by_role("button", name="Trả lương").click()
        s.page.get_by_role("dialog").get_by_role(
            "button", name="Lập chứng từ trả lương"
        ).click()
        s.page.wait_for_load_state("networkidle")
        expect(s.page.locator("main")).to_contain_text("334")
        expect(s.page.locator("main")).to_contain_text("Nháp")
