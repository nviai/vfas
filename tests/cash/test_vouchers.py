"""Cash receipts/payments on the double-entry engine.

v1 section 8 scenario, S04a/S04b books, Excel, PDF, permissions.
"""

from __future__ import annotations

import io
import re
from datetime import date

import pytest
from openpyxl import load_workbook

from app.core.errors import BusinessError, NegativeBalanceError
from app.core.money import amount_in_words
from app.core.registry import ACCOUNTANT, CASHIER, CHIEF, VIEWER
from app.features.cash import books, excel, printing
from app.features.cash import vouchers as mv
from app.features.journal import locks, queries
from app.features.journal import service as js
from app.features.org.models import Company
from tests.conftest import ROLE_EMAILS


def voucher(
    db,
    ctx,
    kind,
    d,
    amount,
    account,
    description="Nghiệp vụ test",
    contact="Nguyễn Văn A",
    post=True,
    **extra,
):
    company = db.get(Company, ctx.company_id)
    form = {
        "date": d.isoformat(),
        "description": description,
        "contact_name": contact,
        "account": account,
        "amount": str(amount),
        **{k: str(v) for k, v in extra.items()},
    }
    v, errors = mv.parse_form(form, kind, company=company, allow_future=True)
    assert v is not None, errors
    e = mv.save(db, ctx, v)
    if post:
        js.post_entry(db, ctx, e)
    return e


@pytest.fixture()
def book(db, ctx, company):
    """Standard data of v1 section 8 (zero opening balance).

    PT 10M 12/01/2023, PC 3M 01/03/2023, PT 5M 05/06/2024.
    """
    company.number_format = "{prefix}{seq:04d}"
    db.commit()
    return {
        "PT0001": voucher(
            db, ctx, "PT", date(2023, 1, 12), 10_000_000, "4111", "Góp vốn"
        ),
        "PC0001": voucher(
            db, ctx, "PC", date(2023, 3, 1), 3_000_000, "6422", "Chi phí"
        ),
        "PT0002": voucher(
            db, ctx, "PT", date(2024, 6, 5), 5_000_000, "5113", "Thu dịch vụ"
        ),
    }


def test_section8_scenarios(db, ctx, book, company):
    assert [e.number for e in book.values()] == ["PT0001", "PC0001", "PT0002"]
    assert queries.balance(db, company.id, "1111", exact=True) == 12_000_000
    # Years 2023 / 2024
    s23 = books.money_book(
        db, company.id, date(2023, 1, 12), date(2023, 12, 31)
    )
    assert (s23.opening, s23.total_in, s23.total_out, s23.closing) == (
        0,
        10_000_000,
        3_000_000,
        7_000_000,
    )
    s24 = books.money_book(
        db, company.id, date(2024, 1, 1), date(2024, 12, 31)
    )
    assert (s24.opening, s24.closing) == (7_000_000, 12_000_000)
    # PC 15M on 10/02/2026 is rejected (only 12M left)
    with pytest.raises(NegativeBalanceError):
        voucher(db, ctx, "PC", date(2026, 2, 10), 15_000_000, "6422")
    # PC 8M on 01/04/2023 is rejected retroactively: row 3, -1M left
    with pytest.raises(NegativeBalanceError) as ei:
        voucher(db, ctx, "PC", date(2023, 4, 1), 8_000_000, "6422")
    assert ei.value.line_no == 3 and ei.value.balance == -1_000_000
    assert "01/04/2023" in ei.value.message
    ok = voucher(db, ctx, "PC", date(2023, 4, 1), 2_000_000, "6422")
    assert ok.status == "posted" and ok.number == "PC0002"


def test_s04a_march_2023_and_detail(db, ctx, book, company):
    voucher(
        db,
        ctx,
        "PT",
        date(2023, 3, 1),
        1_500_000,
        "131",
        partner_id=_customer(db, ctx).id,
    )
    b = books.money_book(db, company.id, date(2023, 3, 1), date(2023, 3, 31))
    assert b.opening == 10_000_000
    assert [
        (r.payment_no or r.receipt_no, r.receipt, r.payment, r.balance)
        for r in b.rows
    ] == [
        ("PC0001", 0, 3_000_000, 7_000_000),
        ("PT0003", 1_500_000, 0, 8_500_000),
    ]
    d = books.money_book(
        db, company.id, date(2023, 3, 1), date(2023, 3, 31), form="s04b"
    )
    assert [r.contra for r in d.rows] == ["6422", "131"]


def _customer(db, ctx):
    from app.features.partners.service import PartnerInput, create

    return create(db, ctx, PartnerInput(name="Công ty ABC", is_customer=True))


def test_multi_line_vat_and_partner_required(db, ctx, company):
    voucher(db, ctx, "PT", date(2024, 1, 2), 50_000_000, "4111")
    company_obj = db.get(Company, ctx.company_id)
    form = {
        "date": "2024-01-05",
        "description": "Mua văn phòng phẩm",
        "contact_name": "Cửa hàng X",
        "mode": "detail",
        "lines-count": "2",
        "lines-0-debit_account": "6422",
        "lines-0-amount": "1.000.000",
        "lines-1-debit_account": "6421",
        "lines-1-amount": "500k",
        "lines-1-description": "Quà tặng KH",
        "vat_on": "1",
        "vat_rate": "10",
        "vat_invoice_no": "0000123",
        "vat_partner_tax_code": "0101243150",
    }
    v, errors = mv.parse_form(form, "PC", company=company_obj)
    assert not errors and v.total == 1_650_000 and v.vat.vat_amount == 150_000
    e = mv.save(db, ctx, v)
    js.post_entry(db, ctx, e)
    assert e.total == 1_650_000
    assert {(ln.debit_account, ln.amount) for ln in e.lines} == {
        ("6422", 1_000_000),
        ("6421", 500_000),
        ("1331", 150_000),
    }
    assert (
        e.vat_records[0].direction == "in"
        and e.vat_records[0].goods_amount == 1_500_000
    )
    # TK 331 requires a party when posting
    e2 = voucher(db, ctx, "PC", date(2024, 1, 6), 100_000, "331", post=False)
    with pytest.raises(BusinessError) as ei:
        js.post_entry(db, ctx, e2)
    assert "đối tượng" in ei.value.message


def test_parse_form_errors(db, company):
    v, errors = mv.parse_form(
        {"date": "2099-01-01", "amount": "abc"}, "PT", company=company
    )
    assert v is None
    assert {"date", "description", "contact_name", "amount", "account"} <= set(
        errors
    )
    v, errors = mv.parse_form(
        {
            "date": "2023-01-11",
            "description": "x",
            "contact_name": "y",
            "account": "711",
            "amount": "1000",
        },
        "PT",
        company=company,
    )
    assert "date" in errors  # before the start date


def test_excel_formulas_recompute(db, ctx, book, company):
    b = books.money_book(db, company.id, date(2023, 1, 1), date(2024, 12, 31))
    data = excel.book_xlsx(
        company,
        b,
        "Từ ngày 01/01/2023 đến ngày 31/12/2024",
        subtitle="Loại quỹ: Tiền Việt Nam",
    )
    ws = load_workbook(io.BytesIO(data))["S04a-DNN"]
    assert ws["H12"].value == 0
    values = {"H12": 0}
    r = 13
    while ws.cell(r, 5).value not in ("Cộng phát sinh trong kỳ", None):
        formula = ws.cell(r, 8).value
        m = re.fullmatch(r"=H(\d+)\+F(\d+)-G(\d+)", formula)
        assert m, formula
        values[f"H{r}"] = (
            values[f"H{m.group(1)}"]
            + (ws.cell(r, 6).value or 0)
            - (ws.cell(r, 7).value or 0)
        )
        assert values[f"H{r}"] == b.rows[r - 13].balance
        r += 1
    assert values[f"H{r - 1}"] == 12_000_000
    assert ws.cell(r, 6).value.startswith("=SUM(F13:")


def test_pdfs(db, ctx, book, company):
    from pypdf import PdfReader

    e = book["PT0001"]
    pdf = printing.voucher_pdf(db, company, e, "Kế toán")
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "PHIẾU THU" in text and "PT0001" in text
    assert amount_in_words(10_000_000).split()[0].capitalize() in text
    b = books.money_book(db, company.id, date(2023, 1, 1), date(2024, 12, 31))
    pdf = printing.book_pdf(
        company,
        b,
        "Từ ngày 01/01/2023 đến ngày 31/12/2024",
        ["Loại quỹ: Tiền Việt Nam"],
    )
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "SỔ QUỸ TIỀN MẶT" in text and "Sổ này có 01 trang" in text


def test_locked_period_cancel_and_adjust(db, ctx, book, company):
    for m in (1, 2, 3):
        locks.lock_period(db, ctx, 2023, m)
    with pytest.raises(BusinessError):
        js.cancel_entry(db, ctx, book["PC0001"], "Nhầm")
    adj = js.reverse_entry(
        db, ctx, book["PC0001"], date(2024, 7, 1), "Chi nhầm"
    )
    js.post_entry(db, ctx, adj)
    assert queries.balance(db, company.id, "1111", exact=True) == 15_000_000
    b = books.money_book(db, company.id, date(2024, 7, 1), date(2024, 7, 31))
    assert b.rows[0].note.startswith("Điều chỉnh cho PC0001")


# ---------------------------------------------------------------- web


def test_web_flow_cashier_quick_form(web, users, db):
    web.login(ROLE_EMAILS[CASHIER])
    r = web.get("/cash/new?kind=PT")
    assert r.status_code == 200 and "Mục đích thu" in r.text
    r = web.post(
        "/cash",
        {
            "kind": "PT",
            "mode": "quick",
            "date": "2024-03-01",
            "doc_date": "2024-03-01",
            "contact_name": "Trần B",
            "description": "Nhận vốn góp",
            "account": "4111",
            "amount": "20tr",
            "action": "post",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text[:500]
    detail = web.get(r.headers["location"])
    assert "PT2024-0001" in detail.text and "20.000.000" in detail.text
    r = web.post(
        "/cash",
        {
            "kind": "PC",
            "mode": "quick",
            "date": "2024-03-02",
            "contact_name": "C",
            "description": "Chi quá tồn",
            "account": "6422",
            "amount": "25tr",
            "action": "post",
        },
        follow_redirects=False,
    )
    # Saved as a draft but cannot be posted (negative cash): the form is
    # shown again with the error and the data kept.
    assert r.status_code == 422
    assert "chưa ghi sổ được" in r.text and "Chi quá tồn" in r.text
    from sqlalchemy import select

    from app.features.journal.models import JournalEntry

    draft = db.scalar(
        select(JournalEntry).where(JournalEntry.description == "Chi quá tồn")
    )
    assert draft is not None and draft.status == "draft"
    assert web.get(draft.url).status_code == 200
    assert web.get("/cash").status_code == 200
    assert (
        web.get(
            "/cash?kind=PC&status=draft",
            headers={"HX-Request": "true", "HX-Target": "cash-results"},
        ).status_code
        == 200
    )
    # The cashier cannot cancel
    pt = db.scalar(
        select(JournalEntry).where(JournalEntry.number == "PT2024-0001")
    )
    assert (
        web.post(f"/entries/{pt.id}/cancel", {"reason": "x"}).status_code
        == 403
    )


def test_web_pages_render(web, users, ctx, db):
    voucher(db, ctx, "PT", date(2024, 1, 2), 5_000_000, "4111")
    e = voucher(db, ctx, "PC", date(2024, 1, 3), 1_000_000, "6422", post=False)
    web.login(ROLE_EMAILS[CHIEF])
    pages = [
        "/cash",
        "/cash/new?kind=PC",
        "/cash/new?kind=PT&mode=detail",
        f"/cash/{e.id}",
        f"/cash/{e.id}/edit",
        "/cash/book",
        "/cash/book?form=s04b&mode=year&year=2024",
        "/cash/book?mode=quarter&quarter=1&year=2024",
        "/cash/counts",
        "/cash/counts/new",
        "/cash/purposes",
        "/cash/import",
        f"/cash/new?copy={e.id}",
    ]
    for url in pages:
        assert web.get(url).status_code == 200, url
    for url in (
        "/cash/book.pdf?mode=year&year=2024",
        "/cash/book.xlsx?form=s04b&mode=year&year=2024",
        f"/cash/{e.id}/print.pdf",
        "/cash/import/template.xlsx",
    ):
        r = web.get(url)
        assert r.status_code == 200 and len(r.content) > 1000, url
    assert web.get("/cash/words?amount=1500000").text.startswith(
        "Một triệu năm trăm nghìn"
    )
    web.get("/lang/en")
    for url in pages:
        assert web.get(url).status_code == 200, url


def test_permissions(web, users):
    web.login(ROLE_EMAILS[VIEWER])
    assert web.get("/cash").status_code == 200
    assert web.get("/cash/new").status_code == 403
    assert web.get("/cash/book.pdf").status_code == 200
    web.login(ROLE_EMAILS[CASHIER])
    assert web.get("/cash/purposes").status_code == 403
    assert web.get("/cash/import").status_code == 403
    web.login(ROLE_EMAILS[ACCOUNTANT])
    assert web.get("/cash/purposes").status_code == 200
