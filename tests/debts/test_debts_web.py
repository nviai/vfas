"""Receivables/payables web.

Summary, aging, party ledger, reconciliation statement PDF, Excel.
"""

import io

from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.registry import ACCOUNTANT, CASHIER, VIEWER
from tests.conftest import ROLE_EMAILS
from tests.debts.test_debts import data  # noqa: F401  (fixture)


def test_debts_pages(web, users, data):  # noqa: F811
    a, b, s = data
    web.login(ROLE_EMAILS[ACCOUNTANT])
    r = web.get("/debts/receivables?mode=year&year=2026")
    assert (
        r.status_code == 200 and "Khách A" in r.text and "4.500.000" in r.text
    )
    assert web.get("/debts/payables").status_code == 200
    r = web.get("/debts/receivables/aging?as_of=2026-10-14")
    assert r.status_code == 200 and "4.500.000" in r.text
    assert web.get("/debts/payables/aging").status_code == 200
    r = web.get(f"/debts/partner/{a.id}?kind=receivable&mode=year&year=2026")
    assert r.status_code == 200 and "Khách A" in r.text
    assert web.get(f"/debts/partner/{s.id}").status_code == 200
    x = web.get("/debts/receivables.xlsx?mode=year&year=2026")
    assert (
        x.status_code == 200
        and load_workbook(io.BytesIO(x.content)).active.max_row > 8
    )
    assert web.get("/debts/receivables/aging.xlsx").status_code == 200
    assert (
        web.get(
            f"/debts/partner/{a.id}/ledger.xlsx?kind=receivable&mode=year&year=2026"
        ).status_code
        == 200
    )
    pdf = web.get(
        f"/debts/partner/{a.id}/reconciliation.pdf?kind=receivable&mode=year&year=2026"
    )
    assert pdf.status_code == 200
    text = "".join(
        p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages
    )
    assert "BIÊN BẢN ĐỐI CHIẾU CÔNG NỢ" in text and "4.500.000" in text
    web.get("/lang/en")
    assert web.get("/debts/receivables").status_code == 200
    assert web.get(f"/debts/partner/{a.id}").status_code == 200


def test_debts_permissions(web, users):
    web.login(ROLE_EMAILS[VIEWER])
    assert web.get("/debts/receivables").status_code == 200
    web.login(ROLE_EMAILS[CASHIER])
    assert web.get("/debts/receivables").status_code == 403
