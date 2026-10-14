"""Sales and purchases web.

Grid form, save draft / post, detail, list, Excel export,
permissions.
"""

import io
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.registry import ACCOUNTANT, CASHIER, CHIEF, VIEWER
from app.features.journal.models import JournalEntry
from app.features.partners import service as ps
from app.features.purchases.models import PurchaseDoc
from app.features.sales.models import SalesDoc
from tests.conftest import ROLE_EMAILS

FIX = Path(__file__).parent.parent / "purchases" / "fixtures"


def _grid(lines):
    data = {"lines-count": str(len(lines))}
    for i, ln in enumerate(lines):
        for k, v in ln.items():
            data[f"lines-{i}-{k}"] = str(v)
    return data


def test_sales_create_post_detail_list_export(web, users, db, ctx):
    c = ps.create(
        db,
        ctx,
        ps.PartnerInput(
            name="Khách Web", is_customer=True, tax_code="0101243150"
        ),
    )
    web.login(ROLE_EMAILS[ACCOUNTANT])
    assert web.get("/sales/new").status_code == 200
    data = {
        "kind": "sale",
        "date": "2026-03-05",
        "partner_id": str(c.id),
        "invoice_series": "c26taa",
        "invoice_no": "15",
        "payment_method": "credit",
        "action": "post",
        **_grid(
            [
                {
                    "description": "Phí dịch vụ",
                    "quantity": "2",
                    "unit_price": "1500000",
                    "amount": "3000000",
                    "vat_rate": "10",
                    "vat_amount": "300000",
                    "account": "",
                },
                {
                    "description": "Hàng",
                    "amount": "1.000.000",
                    "vat_rate": "8",
                },
            ]
        ),
    }
    r = web.post("/sales", data, follow_redirects=False)
    assert r.status_code == 303, r.text[:400]
    doc = db.scalars(select(SalesDoc)).one()
    db.refresh(doc)
    entry = db.get(JournalEntry, doc.entry_id)
    assert (
        entry.status == "posted"
        and doc.total_amount == 4_380_000
        and doc.invoice_series == "C26TAA"
    )
    assert r.headers["location"] == f"/sales/{entry.id}"
    page = web.get(f"/sales/{entry.id}")
    assert (
        page.status_code == 200
        and "BH2026-0001" in page.text
        and "4.380.000" in page.text
    )
    assert entry.number in web.get("/sales").text
    assert web.get("/sales?status=draft").status_code == 200
    x = web.get("/sales/export.xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.active.max_row >= 10
    # Create a return from the original voucher
    r = web.get(f"/sales/new?source={entry.id}")
    assert r.status_code == 200 and "Giảm giá, trả lại" in r.text
    assert web.get("/sales/returns").status_code == 200
    web.get("/lang/en")
    assert web.get(f"/sales/{entry.id}").status_code == 200
    assert web.get("/sales/new").status_code == 200


def test_sales_form_errors_and_draft_edit(web, users, db):
    web.login(ROLE_EMAILS[CHIEF])
    web.get("/sales/new")
    r = web.post(
        "/sales",
        {
            "kind": "sale",
            "date": "",
            "payment_method": "credit",
            "lines-count": "0",
        },
    )
    assert r.status_code == 422 and "Nhập ngày" in r.text
    data = {
        "kind": "sale",
        "date": "2026-03-05",
        "partner_id": "Khách hàng gõ tay",
        "payment_method": "cash",
        "action": "save",
        **_grid(
            [{"description": "Bán lẻ", "amount": "500000", "vat_rate": "10"}]
        ),
    }
    r = web.post("/sales", data, follow_redirects=False)
    assert r.status_code == 303
    doc = db.scalars(select(SalesDoc)).one()
    assert doc.partner_name == "Khách hàng gõ tay"
    assert web.get(f"/sales/{doc.entry_id}/edit").status_code == 200
    data["lines-0-amount"] = "700000"
    data["lines-0-vat_amount"] = ""
    r = web.post(f"/sales/{doc.entry_id}", data, follow_redirects=False)
    assert r.status_code == 303
    db.refresh(doc)
    assert doc.total_amount == 770_000
    # Does posting fail (no cash yet? no - a cash sale cannot make cash
    # negative) -> it posts
    r = web.post(f"/entries/{doc.entry_id}/post", follow_redirects=False)
    assert (
        r.status_code == 303
        and r.headers["location"] == f"/sales/{doc.entry_id}"
    )


def test_purchase_web_and_permissions(web, users, db):
    web.login(ROLE_EMAILS[ACCOUNTANT])
    data = {
        "kind": "purchase",
        "date": "2026-03-05",
        "partner_id": "Nhà cung cấp mới",
        "tax_code": "0301234567",
        "payment_method": "credit",
        "action": "post",
        **_grid(
            [
                {
                    "description": "Máy in",
                    "amount": "8000000",
                    "vat_rate": "10",
                    "account": "153",
                    "deductible": "true",
                },
                {
                    "description": "Tiếp khách",
                    "amount": "1000000",
                    "vat_rate": "10",
                    "account": "6422",
                    "deductible": "false",
                },
            ]
        ),
    }
    r = web.post("/purchases", data, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    doc = db.scalars(select(PurchaseDoc)).one()
    entry = db.get(JournalEntry, doc.entry_id)
    assert entry.status == "posted"
    assert sorted((ln.debit_account, ln.amount) for ln in entry.lines) == [
        ("1331", 800_000),
        ("153", 8_000_000),
        ("6422", 1_100_000),
    ]
    assert web.get(f"/purchases/{entry.id}").status_code == 200
    assert (
        web.get("/purchases").status_code == 200
        and web.get("/purchases/returns").status_code == 200
    )
    web.login(ROLE_EMAILS[CASHIER])
    assert web.get("/purchases").status_code == 403
    assert web.get("/sales/new").status_code == 403
    web.login(ROLE_EMAILS[VIEWER])
    assert web.get("/purchases").status_code == 200
    assert web.get("/purchases/new").status_code == 403


def test_einvoice_web_flow(web, users, db):
    web.login(ROLE_EMAILS[ACCOUNTANT])
    assert web.get("/einvoices").status_code == 200
    files = [
        (
            "files",
            ("hd1.xml", (FIX / "hd_dau_vao_1.xml").read_bytes(), "text/xml"),
        ),
        (
            "files",
            ("hd2.xml", (FIX / "hd_dau_ra_1.xml").read_bytes(), "text/xml"),
        ),
    ]
    r = web.post("/einvoices", files=files)
    assert (
        r.status_code == 200 and "00000457" in r.text and "Sẽ nhập" in r.text
    )
    import re

    token = re.search(r"/einvoices/([0-9a-f]+)/confirm", r.text).group(1)
    keys = re.findall(r'name="keys" value="([^"]+)"', r.text)
    r = web.post(
        f"/einvoices/{token}/confirm",
        {"keys": keys, "purchase_account": "6422"},
        follow_redirects=False,
    )
    assert r.status_code == 303
    assert db.scalar(select(PurchaseDoc.id)) and db.scalar(select(SalesDoc.id))
    _ = date
