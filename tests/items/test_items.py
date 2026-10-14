"""Goods and services catalog.

Add/edit on the web, Excel import/export, permissions.
"""

import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.registry import ACCOUNTANT, CASHIER, VIEWER
from app.features.items.models import Item
from tests.conftest import ROLE_EMAILS


def _xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Mã",
            "Tên",
            "ĐVT",
            "Loại",
            "Thuế suất",
            "Giá bán",
            "Giá mua",
            "TK kho",
            "TK doanh thu",
            "TK chi phí",
        ]
    )
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_items_crud_import_export(web, users, db):
    web.login(ROLE_EMAILS[ACCOUNTANT])
    assert web.get("/items").status_code == 200
    assert web.get("/items/new?kind=service").status_code == 200
    r = web.post(
        "/items",
        {
            "name": "Tư vấn triển khai",
            "kind": "service",
            "unit": "Giờ",
            "vat_rate": "10",
            "sale_price": "500.000",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    it = db.scalars(select(Item)).one()
    assert (
        it.code == "DV0001"
        and it.revenue_account == "5113"
        and int(it.sale_price) == 500_000
    )
    assert web.get(f"/items/{it.id}").status_code == 200
    r = web.post(
        f"/items/{it.id}",
        {
            "name": "Tư vấn triển khai phần mềm",
            "kind": "service",
            "vat_rate": "8",
            "code": "DV0001",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303
    db.refresh(it)
    assert it.vat_rate == "8"
    # Excel import: any error -> no row is imported
    bad = _xlsx(
        [
            [
                "HH0001",
                "Laptop",
                "Chiếc",
                "Hàng hóa",
                "10",
                25_000_000,
                21_000_000,
                "",
                "",
                "",
            ],
            ["X1", "Sai loại", "", "Không biết", "10", 0, 0, "", "", ""],
        ]
    )
    r = web.post(
        "/items/import",
        files={"file": ("a.xlsx", bad, "application/octet-stream")},
    )
    assert (
        r.status_code == 422
        and "Sai loại" not in r.text
        and "không hợp lệ" in r.text
    )
    assert len(db.scalars(select(Item)).all()) == 1
    good = _xlsx(
        [
            [
                "HH0001",
                "Laptop",
                "Chiếc",
                "Hàng hóa",
                "10%",
                25_000_000,
                21_000_000,
                "",
                "",
                "",
            ],
            [
                "",
                "Giấy A4",
                "Ram",
                "Nguyên vật liệu",
                "8",
                "75.000",
                "",
                "",
                "",
                "",
            ],
        ]
    )
    r = web.post(
        "/items/import",
        files={"file": ("a.xlsx", good, "application/octet-stream")},
        follow_redirects=False,
    )
    assert r.status_code == 303
    items = {i.name: i for i in db.scalars(select(Item))}
    assert (
        items["Laptop"].stock_account == "156"
        and items["Giấy A4"].code == "VT0001"
    )
    x = web.get("/items/export.xlsx")
    assert (
        x.status_code == 200
        and load_workbook(io.BytesIO(x.content)).active.max_row >= 10
    )
    assert web.get("/items/import/template.xlsx").status_code == 200
    r = web.get("/lookup/items?q=lap")
    assert r.json()[0]["extra"]["sale_price"].startswith("25000000")
    web.get("/lang/en")
    assert (
        web.get("/items").status_code == 200
        and web.get(f"/items/{it.id}").status_code == 200
    )


def test_items_permissions(web, users):
    web.login(ROLE_EMAILS[VIEWER])
    assert web.get("/items").status_code == 200
    assert web.get("/items/new").status_code == 403
    web.login(ROLE_EMAILS[CASHIER])
    assert web.get("/items").status_code == 200
    web.get("/items")
    assert web.post("/items", {"name": "X"}).status_code == 403
