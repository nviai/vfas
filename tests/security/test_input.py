"""Review C regressions (C-1..C-11): input and import/export files."""

from __future__ import annotations

import csv
import io
import json
import os
import re
import time
import zipfile
from datetime import date
from html.parser import HTMLParser
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.config import APP_DIR
from app.core.errors import BusinessError, NotFoundError
from app.core.registry import CHIEF
from tests.conftest import ROLE_EMAILS

FIX = Path(__file__).parent.parent / "purchases" / "fixtures"
EVIL_FORMULA = '=HYPERLINK("http://x")'


# ---------------------------------------- C-1 XSS via Alpine attributes


class _Attrs(HTMLParser):
    def __init__(self):
        super().__init__()
        self.attrs: list[tuple[str, str, str | None]] = []

    def handle_starttag(self, tag, attrs):
        self.attrs.extend((tag, k, v) for k, v in attrs)


@pytest.mark.parametrize(
    "payload",
    [
        "a x-init=window.__xss=5 b",
        "'><img src=x onerror=1>",
        '" @click="window.__xss=1',
    ],
)
def test_warehouse_page_escapes_user_data_in_attributes(
    web, users, db, ctxs, payload
):
    from app.features.inventory import service as inv

    inv.save_warehouse(
        db,
        ctxs[CHIEF],
        wh=None,
        code="K2",
        name=payload,
        address=payload,
        keeper=payload,
    )
    web.login(ROLE_EMAILS[CHIEF])
    r = web.get("/inventory/warehouses")
    assert r.status_code == 200
    p = _Attrs()
    p.feed(r.text)
    names = {k for _t, k, _v in p.attrs}
    assert "x-init" not in names and "onerror" not in names
    assert not any(
        t == "img" and v == "x" for t, k, v in p.attrs if k == "src"
    )
    # Warehouse data lives only in data-wh (valid JSON); the edit button
    # reads the exact values back.
    blobs = [json.loads(v) for _t, k, v in p.attrs if k == "data-wh"]
    k2 = next(b for b in blobs if b["code"] == "K2")
    assert k2["address"] == payload and k2["name"] == payload
    clicks = [
        v for _t, k, v in p.attrs if k == "@click" and "edit(" in (v or "")
    ]
    assert clicks and all(
        v == "edit(JSON.parse($el.dataset.wh))" for v in clicks
    )


_DANGEROUS_TEMPLATE = [
    # String concatenation into attrs= (the macro returns Markup) with
    # JSON/forceescape/safe: a quote in the data escapes the attribute.
    re.compile(r"attrs=(?:(?!\}\}).)*(?:tojson|forceescape|\|\s*safe\b)"),
    # tojson/forceescape output concatenated with ~ (leaves the safe
    # attribute context).
    re.compile(r"(?:tojson|forceescape)\s*\)?\s*~"),
    re.compile(r"~\s*\(?\s*[\w.\[\]'\"]+\s*\|\s*(?:tojson|forceescape)"),
]


def _dangerous(line: str) -> bool:
    return any(p.search(line) for p in _DANGEROUS_TEMPLATE)


def test_static_scan_templates_no_attribute_concatenation():
    bad_samples = [
        """{{ ui.icon_button("pencil", "Sửa", attrs="@click='edit(" ~ """
        """(w|tojson|forceescape) ~ ")'") }}""",
        """<button @click="edit({{ '' ~ w|tojson }})">""",
        """{{ ui.button("x", attrs="@click='go(" ~ data|tojson ~ ")'") }}""",
    ]
    assert all(_dangerous(s) for s in bad_samples)
    hits = []
    for path in (APP_DIR).rglob("*.html"):
        for n, line in enumerate(
            path.read_text(encoding="utf-8").splitlines(), start=1
        ):
            if _dangerous(line):
                hits.append(
                    f"{path.relative_to(APP_DIR)}:{n}: {line.strip()[:120]}"
                )
    assert not hits, hits


# ------------------------------------------ C-2 Excel formula injection


def _cells(data: bytes):
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            yield from row


def test_reportbook_formula_injection_neutralized():
    from app.core.xlsx import Col, ReportBook

    book = ReportBook()
    ws = book.add_sheet(
        "S", title="T", columns=[Col("Tên", 20), Col("Tiền", 12, "money")]
    )
    book.row(ws, [EVIL_FORMULA, 5])
    book.row(ws, ["=cmd|' /C calc'!A0", 7])
    book.total_row(ws, "Cộng", sum_cols=[1])
    cells = list(_cells(book.to_bytes()))
    evil = [
        c
        for c in cells
        if isinstance(c.value, str)
        and c.value.startswith(("=HYPERLINK", "=cmd"))
    ]
    assert len(evil) == 2 and all(c.data_type == "s" for c in evil)
    sums = [
        c
        for c in cells
        if isinstance(c.value, str) and c.value.startswith("=SUM(")
    ]
    assert sums and all(c.data_type == "f" for c in sums)


def test_exports_keep_user_text_as_text(web, users, db, ctxs):
    from app.features.items import service as items

    items.create(
        db, ctxs[CHIEF], items.ItemInput(name=EVIL_FORMULA, code="HH9")
    )
    web.login(ROLE_EMAILS[CHIEF])
    r = web.post(
        "/cash",
        {
            "kind": "PT",
            "mode": "quick",
            "date": "2024-03-01",
            "doc_date": "2024-03-01",
            "contact_name": EVIL_FORMULA,
            "description": EVIL_FORMULA,
            "account": "4111",
            "amount": "20tr",
            "action": "post",
        },
        follow_redirects=False,
    )
    assert r.status_code == 303, r.text[:300]

    r = web.get("/items/export.xlsx")
    assert r.status_code == 200
    evil = [c for c in _cells(r.content) if c.value == EVIL_FORMULA]
    assert evil and all(c.data_type == "s" for c in evil)

    r = web.get("/cash/book.xlsx?mode=all")
    assert r.status_code == 200
    cells = list(_cells(r.content))
    evil = [c for c in cells if c.value == EVIL_FORMULA]
    assert evil and all(c.data_type == "s" for c in evil)
    formulas = [
        c
        for c in cells
        if isinstance(c.value, str)
        and c.value.startswith("=")
        and c.value != EVIL_FORMULA
    ]
    # The cash balance column and the totals row stay formulas.
    assert formulas and all(c.data_type == "f" for c in formulas)


# -------------------------------------------- C-3 voucher number format


@pytest.mark.parametrize(
    "fmt",
    [
        "{prefix}{seq:>50000000}",
        "{seq.__class__}",
        "{prefix}{year}",
        "{seq}{seq}",
        "A" * 25 + "{seq}",
        "{prefix}{seq:011d}",
        "{seq!r}",
        "<b>{seq}</b>",
    ],
)
def test_number_format_rejected(db, ctx, company, fmt):
    from app.features.journal.service import validate_number_format
    from app.features.org.service import update_company

    assert validate_number_format(fmt)
    with pytest.raises(BusinessError):
        update_company(db, ctx, {"number_format": fmt})
    db.rollback()


def test_number_format_defensive_fallback(db, ctx):
    from app.features.journal.service import (
        format_number,
        validate_number_format,
    )
    from app.features.org.service import update_company_settings

    for ok in (
        "{prefix}{year}-{seq:04d}",
        "{prefix}/{yy}{month}/{seq:06d}",
        "HD.{seq}",
    ):
        assert validate_number_format(ok) is None
    d = date(2026, 3, 5)
    assert (
        format_number("{prefix}{seq:>50000000}", "PT", d, 5) == "PT2026-0005"
    )
    assert (
        format_number("{prefix}{year}-{seq:04d}", "{x}", d, 5) == "CT2026-0005"
    )
    assert (
        len(
            format_number(
                "{prefix}" + "ABCDEFGHIJ" * 2 + "{seq:010d}", "PT", d, 5
            )
        )
        <= 40
    )
    with pytest.raises(BusinessError):
        update_company_settings(db, ctx, "prefixes", {"PT": "{seq:>99999}"})
    db.rollback()


# ------------------------------------------ C-4 import file size limits


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_items_import_row_limit(db, ctx):
    from app.features.items.router import (
        IMPORT_COLUMNS,
        MAX_IMPORT_ROWS,
        import_rows,
    )

    data = _xlsx(
        [IMPORT_COLUMNS]
        + [["", f"Hàng {i}"] for i in range(MAX_IMPORT_ROWS + 2)]
    )
    with pytest.raises(BusinessError, match="quá nhiều dòng"):
        import_rows(db, ctx, data)


def test_items_import_assigns_codes_without_rescanning(db, ctx, monkeypatch):
    from app.features.items import router as items_router
    from app.features.items import service as items

    calls = []
    real = items.next_code
    monkeypatch.setattr(
        items, "next_code", lambda *a, **kw: calls.append(a) or real(*a, **kw)
    )
    rows = [items_router.IMPORT_COLUMNS] + [
        ["", f"Hàng {i}", "Cái", "Hàng hóa"] for i in range(300)
    ]
    rows += [["", f"Dịch vụ {i}", "Giờ", "Dịch vụ"] for i in range(50)]
    rows.append(["", "Thuế sai", "Cái", "Hàng hóa", "0.abc"])
    t = time.perf_counter()
    result = items_router.import_rows(db, ctx, _xlsx(rows))
    # A data error, not a 500.
    assert result["created"] == 0 and any(
        "Dòng 352" in e for e in result["errors"]
    )
    result = items_router.import_rows(db, ctx, _xlsx(rows[:-1]))
    assert result == {"created": 350, "errors": []}, result["errors"][:3]
    # Once per prefix (per import run), not once per row.
    assert len(calls) <= 4
    assert time.perf_counter() - t < 60
    codes = set(
        db.scalars(
            select(items.Item.code).where(
                items.Item.company_id == ctx.company_id
            )
        )
    )
    assert {"HH0001", "HH0300", "DV0001", "DV0050"} <= codes


def test_bank_statement_column_limit():
    from app.features.bank import statements as st

    wide = "\n".join(
        ",".join(str(i) for i in range(500)) for _ in range(3)
    ).encode()
    rows = st.read_rows(wide, "sao-ke.csv")
    assert rows and max(len(r) for r in rows) <= st.MAX_COLS
    rows = st.read_rows(_xlsx([list(range(300))] * 3), "sao-ke.xlsx")
    assert rows and max(len(r) for r in rows) <= st.MAX_COLS
    huge_field = ("a," + "x" * (csv.field_size_limit() + 10) + "\n").encode()
    with pytest.raises(BusinessError):
        st.read_rows(huge_field, "sao-ke.csv")


# ------------------------------------- C-5 huge numbers / corrupt files


def _invoice(**replace) -> bytes:
    xml = (FIX / "hd_dau_vao_1.xml").read_text(encoding="utf-8")
    for old, new in replace.items():
        xml = xml.replace(old, new, 1)
    return xml.encode()


BAD_INVOICES = {
    "sl-1e999999.xml": _invoice(
        **{"<SLuong>20</SLuong>": "<SLuong>1e999999</SLuong>"}
    ),
    "dgia-nan.xml": _invoice(**{"<DGia>75000</DGia>": "<DGia>NaN</DGia>"}),
    "tien-30-so.xml": _invoice(
        **{"<ThTien>1500000</ThTien>": "<ThTien>" + "9" * 30 + "</ThTien>"}
    ),
    "tien-mu.xml": _invoice(
        **{"<TgTTTBSo>2611999</TgTTTBSo>": "<TgTTTBSo>1E+400</TgTTTBSo>"}
    ),
    "tsuat.xml": _invoice(**{"<TSuat>8%</TSuat>": "<TSuat>1e999999%</TSuat>"}),
    "tsuat-nan.xml": _invoice(**{"<TSuat>10%</TSuat>": "<TSuat>NaN%</TSuat>"}),
}


def _zip(
    members: dict[str, bytes],
    *,
    encrypted: bool = False,
    corrupt: bool = False,
) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    out = bytearray(buf.getvalue())
    if encrypted:
        # Set the "encrypted" flag in the local header (+6) and the
        # central directory (+8) -> zipfile asks for a password.
        for sig, off in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
            i = out.find(sig)
            while i >= 0:
                out[i + off] |= 0x1
                i = out.find(sig, i + 4)
    out = bytes(out)
    if corrupt:  # change one content byte -> CRC mismatch
        i = out.index(b"<HDon>")
        out = out[:i] + b"X" + out[i + 1 :]
    return out


def test_einvoice_parser_never_crashes():
    from app.features.purchases.einvoice import read_files

    good = (FIX / "hd_dau_vao_1.xml").read_bytes()
    files = list(BAD_INVOICES.items()) + [
        ("hong.zip", _zip({"a.xml": good}, corrupt=True)),
        ("mat-khau.zip", _zip({"a.xml": good}, encrypted=True)),
        ("vo.zip", b"PK\x03\x04" + b"\x00" * 50),
        (
            "trong-zip.zip",
            _zip({"x/" + k: v for k, v in BAD_INVOICES.items()}),
        ),
    ]
    invoices, errors = read_files(files)
    assert all(
        abs(ln.amount) <= 10**15 for inv in invoices for ln in inv.lines
    )
    tsuat = [inv for inv in invoices if inv.filename.startswith("tsuat")]
    assert len(tsuat) == 4 and all(
        any(ln.vat_rate == "KHAC" for ln in inv.lines) for inv in tsuat
    )
    for name in (
        "sl-1e999999.xml",
        "dgia-nan.xml",
        "tien-30-so.xml",
        "tien-mu.xml",
        "hong.zip",
        "mat-khau.zip",
        "vo.zip",
    ):
        assert any(name in e for e in errors), (name, errors)


def test_einvoice_upload_bad_numbers_not_500(web, users):
    web.login(ROLE_EMAILS[CHIEF])
    files = [
        ("files", (name, data, "text/xml"))
        for name, data in BAD_INVOICES.items()
    ]
    files.append(
        (
            "files",
            (
                "hong.zip",
                _zip({"a.xml": BAD_INVOICES["dgia-nan.xml"]}, corrupt=True),
                "application/zip",
            ),
        )
    )
    r = web.post("/einvoices", files=files, follow_redirects=False)
    assert r.status_code in (200, 303), r.status_code


def test_numeric_parsers_bounded():
    from decimal import Decimal

    from app.core.money import MoneyError, parse_vnd
    from app.features.bank.statements import parse_amount
    from app.features.cash.importer import _amount, _balance
    from app.features.partners.router import _input
    from app.features.sales.common import (
        parse_decimal,
        parse_money,
        parse_vn_decimal,
    )

    for bad in (
        "9" * 30,
        "1" * 5000,
        "9.999.999.999.999.999",
        "2000000tỷ",
        10**16,
    ):
        with pytest.raises(MoneyError):
            parse_vnd(bad)
    assert parse_vnd("999.999.999.999.999") == 999_999_999_999_999
    for bad in (
        "9" * 30,
        "1e999",
        "1E+400",
        "NaN",
        "-Infinity",
        "1" * 50,
        Decimal("1e999"),
    ):
        with pytest.raises(ValueError):
            parse_decimal(bad)
        with pytest.raises(ValueError):
            parse_vn_decimal(bad) if not isinstance(
                bad, Decimal
            ) else parse_decimal(bad)
    assert parse_decimal("1e-999999") == 0 and parse_decimal(
        "1.234,5"
    ) == Decimal("1234.5")
    with pytest.raises(MoneyError):
        parse_money("9" * 30)
    for bad in (
        "9" * 30,
        float("inf"),
        float("nan"),
        Decimal("NaN"),
        Decimal("1e999"),
        10**20,
    ):
        assert parse_amount(bad) is None
    assert parse_amount("1.000.000") == 1_000_000
    for bad in (
        "9" * 30,
        float("inf"),
        Decimal("1e999"),
        Decimal("NaN"),
        10**20,
    ):
        value, issue = _amount(bad)
        assert value is None and issue is not None
        assert _balance(bad) is None
    assert (
        _input(
            {"name": "A", "payment_terms_days": "99999999"}
        ).payment_terms_days
        == 3650
    )
    assert (
        _input(
            {"name": "A", "payment_terms_days": "9" * 5000}
        ).payment_terms_days
        == 0
    )


@pytest.mark.parametrize(
    "field,value",
    [
        ("quantity", "9" * 30),
        ("quantity", "1e999"),
        ("unit_price", "NaN"),
        ("amount", "9" * 30),
        ("quantity", "1000000000000"),
    ],
)
def test_sales_huge_numbers_friendly_error(web, users, db, ctx, field, value):
    from app.features.partners import service as ps
    from app.features.sales.models import SalesDoc

    c = ps.create(
        db, ctx, ps.PartnerInput(name="Khách số lớn", is_customer=True)
    )
    web.login(ROLE_EMAILS[CHIEF])
    line = {
        "description": "Hàng",
        "quantity": "2",
        "unit_price": "1000",
        "amount": "",
        "vat_rate": "10",
    }
    line[field] = value
    if field == "quantity" and value == "1000000000000":
        line["unit_price"] = "1000000000000"  # line amount 10^24
    data = {
        "kind": "sale",
        "date": "2026-03-05",
        "partner_id": str(c.id),
        "payment_method": "credit",
        "action": "post",
        "lines-count": "1",
        **{f"lines-0-{k}": v for k, v in line.items()},
    }
    r = web.post("/sales", data, follow_redirects=False)
    assert r.status_code < 500
    assert db.scalar(select(SalesDoc.id)) is None


# ------------------------ C-6 Unicode file names in Content-Disposition


def test_unicode_download_filenames(web, users, db, ctx):
    from app.core.xlsx import content_disposition
    from app.features.partners import service as ps

    h = content_disposition('so "chi" tiết\r\nX-Evil: 1\\KHĐ01.xlsx')
    assert "\r" not in h and "\n" not in h and h.isascii()
    assert 'filename="' in h and "filename*=UTF-8''" in h
    p = ps.create(
        db,
        ctx,
        ps.PartnerInput(name="Khách Đ", code="KHĐ01", is_customer=True),
    )
    web.login(ROLE_EMAILS[CHIEF])
    r = web.get(f"/debts/partner/{p.id}/ledger.xlsx")
    assert (
        r.status_code == 200
        and "KH%C4%9001" in r.headers["content-disposition"]
    )
    r = web.get(f"/debts/partner/{p.id}/reconciliation.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert "KH%C4%9001" in r.headers["content-disposition"]


# ----------------------------------- C-7 control characters in markdown


def test_markdown_control_chars():
    from app.features.ai.markdown import render

    assert render("\x000\x00") == "<p>0</p>"
    out = render("`code` \x000\x00 \x001\x00 **b**\x07")
    assert "\x00" not in out and "\x07" not in out and out.count("<code>") == 1
    assert render("\x00\x01") == ""


# ------------------------------------------------ C-8 temp file cleanup


def _age(path: Path, hours: float) -> None:
    t = time.time() - hours * 3600
    os.utime(path, (t, t))


def test_cleanup_imports(tmp_path):
    from app.core.tempfiles import cleanup_imports

    old_file, new_file = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    nested = tmp_path / "einvoices" / "1" / "tok"
    nested.mkdir(parents=True)
    empty_old = tmp_path / "empty"
    empty_old.mkdir()
    for f in (old_file, new_file, nested / "0000.xml"):
        f.write_bytes(b"x")
    _age(old_file, 30)
    _age(nested / "0000.xml", 30)
    _age(empty_old, 30)
    assert cleanup_imports(24, base=tmp_path) >= 3
    assert (
        new_file.exists()
        and not old_file.exists()
        and not (nested / "0000.xml").exists()
    )
    assert not empty_old.exists()


def test_new_import_session_cleans_old_files(settings, company, monkeypatch):
    from app.core import tempfiles
    from app.features.bank import statements as st

    old = settings.imports_dir / "bank-old.bin"
    old.write_bytes(b"x")
    _age(old, 48)
    monkeypatch.setattr(tempfiles, "_last_run", 0.0)
    st.save_upload(company.id, 1, "sao-ke.csv", b"a,b\n1,2\n")
    assert not old.exists()


# ----------------------- C-9 WeasyPrint reads only data: and app/static


def test_pdf_fetcher_blocks_external_and_local_files(tmp_path):
    from app.core.pdf import _html, _safe_fetcher

    f = _safe_fetcher()
    assert _html("<p>x</p>").url_fetcher is f
    secret = tmp_path / "secret.css"
    secret.write_text("body{}")
    for url in (
        "http://169.254.169.254/latest/meta-data/",
        "https://example.com/a.png",
        "ftp://x/y",
        f"file://{secret}",
        "file:///etc/passwd",
        f"file://{APP_DIR / 'static'}/../core/config.py",
    ):
        with pytest.raises(ValueError):
            f.fetch(url)
    static_file = next(
        p for p in (APP_DIR / "static").rglob("*") if p.is_file()
    )
    f.fetch(static_file.resolve().as_uri())
    f.fetch("data:text/plain,abc")


# ----------------------------- C-10 cancelling another company's import


def test_discard_upload_checks_company(settings, company):
    from app.features.bank import statements as st
    from app.features.cash import importer

    token = importer.save_upload(
        company.id, "so-quy.xlsx", b"PK\x03\x04data", None
    )
    with pytest.raises(NotFoundError):
        importer.discard_upload(token, company.id + 1)
    assert importer.load_upload(company.id, token)
    importer.discard_upload(token, company.id)
    with pytest.raises(NotFoundError):
        importer.load_upload(company.id, token)

    token = st.save_upload(company.id, 1, "sao-ke.csv", b"a,b\n")
    with pytest.raises(NotFoundError):
        st.discard_upload(token, company.id + 1)
    assert st.load_upload(company.id, token)
    st.discard_upload(token, company.id)
    with pytest.raises(NotFoundError):
        st.load_upload(company.id, token)


# -------------------------------------------------- C-11 LIKE wildcards


def test_like_wildcards_escaped(db, ctx, company):
    from app.core.db import like_escape
    from app.features.accounts.service import list_accounts
    from app.features.items import service as items

    assert like_escape("50%_a\\") == "50\\%\\_a\\\\"
    items.create(db, ctx, items.ItemInput(name="Giảm 50% phí", code="HH1"))
    items.create(db, ctx, items.ItemInput(name="Bút bi", code="HH2"))
    # Only names that really contain "%".
    assert [i.code for i in items.search(db, company.id, "%")] == ["HH1"]
    assert [i.code for i in items.search(db, company.id, "50%")] == ["HH1"]
    assert items.search(db, company.id, "_") == []
    assert items.search(db, company.id, "b_t") == []
    assert list_accounts(db, company.id, q="%") == []
    assert list_accounts(db, company.id, q="_") == []
    assert [a.code for a in list_accounts(db, company.id, q="1111")] == [
        "1111"
    ]
