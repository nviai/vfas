"""Materials, goods and services catalog.

List, create/edit, and Excel import/export.
"""

from __future__ import annotations

import io
from decimal import Decimal

from fastapi import APIRouter, Depends, Request
from sqlalchemy.orm import Session

from app.core.errors import BusinessError
from app.core.security import CurrentUser, require
from app.core.web import (
    form_data,
    get_ctx,
    get_tenant_db,
    is_htmx,
    redirect,
    render,
)
from app.core.xlsx import Col, ReportBook, save_workbook, xlsx_response
from app.features.items import service
from app.features.items.models import ITEM_KINDS, VAT_RATES
from app.features.items.service import ItemInput
from app.features.org.models import Company
from app.features.sales.common import parse_vn_decimal as parse_decimal

MAX_IMPORT_ROWS = 5000
router = APIRouter()
PAGE = 50
IMPORT_COLUMNS = [
    "Mã",
    "Tên",
    "ĐVT",
    "Loại",
    "Thuế suất",
    "Giá bán",
    "Giá mua",
    "TK kho",
    "TK doanh thu",
    "TK chi phí",
]
KIND_BY_LABEL = {v.lower(): k for k, v in ITEM_KINDS.items()} | {
    k: k for k in ITEM_KINDS
}


def _input(form: dict) -> ItemInput:
    try:
        sale = parse_decimal(form.get("sale_price")) or Decimal(0)
        purchase = parse_decimal(form.get("purchase_price")) or Decimal(0)
    except ValueError as e:
        raise BusinessError("Đơn giá không hợp lệ") from e
    return ItemInput(
        name=form.get("name") or "",
        code=form.get("code") or "",
        unit=(form.get("unit") or "").strip(),
        kind=form.get("kind") or "goods",
        stock_account=form.get("stock_account") or None,
        revenue_account=form.get("revenue_account") or None,
        cogs_account=form.get("cogs_account") or "632",
        expense_account=form.get("expense_account") or None,
        vat_rate=form.get("vat_rate") or "10",
        sale_price=sale,
        purchase_price=purchase,
        note=form.get("note") or "",
        is_active=form.get("is_active", "1") == "1",
    )


def _ctx(extra: dict) -> dict:
    return {"kinds": ITEM_KINDS, "vat_rates": VAT_RATES, **extra}


@router.get("/items")
def items_page(
    request: Request,
    q: str = "",
    kind: str = "",
    show: str = "active",
    page: int = 1,
    user: CurrentUser = Depends(require("catalog.view")),
    db: Session = Depends(get_tenant_db),
):
    page = max(1, page)
    rows = service.search(
        db,
        user.company_id,
        q,
        kind=kind or None,
        active=None if show == "all" else True,
        limit=PAGE + 1,
        offset=(page - 1) * PAGE,
    )
    ctx = _ctx(
        {
            "items": rows[:PAGE],
            "has_next": len(rows) > PAGE,
            "page": page,
            "q": q,
            "kind": kind,
            "show": show,
            "breadcrumbs": [("Danh mục", None), ("Hàng hóa, dịch vụ", None)],
        }
    )
    if is_htmx(request) and request.headers.get("hx-target") == "item-results":
        return render(request, "items/_rows.html", ctx)
    return render(request, "items/index.html", ctx)


@router.get("/items/new")
def item_new(
    request: Request,
    kind: str = "goods",
    _u: CurrentUser = Depends(require("catalog.create")),
):
    kind = kind if kind in ITEM_KINDS else "goods"
    return render(
        request,
        "items/form.html",
        _ctx(
            {
                "it": None,
                "form": {"kind": kind, "vat_rate": "10"},
                "errors": {},
                "breadcrumbs": [
                    ("Hàng hóa, dịch vụ", "/items"),
                    ("Thêm mới", None),
                ],
            }
        ),
    )


@router.post("/items")
def item_create(
    request: Request,
    form: dict = Depends(form_data),
    _u: CurrentUser = Depends(require("catalog.create")),
    db: Session = Depends(get_tenant_db),
):
    try:
        it = service.create(db, get_ctx(request), _input(form))
    except BusinessError as e:
        db.rollback()
        return render(
            request,
            "items/form.html",
            _ctx(
                {
                    "it": None,
                    "form": form,
                    "errors": {"name": e.message},
                    "breadcrumbs": [
                        ("Hàng hóa, dịch vụ", "/items"),
                        ("Thêm mới", None),
                    ],
                }
            ),
            status_code=422,
        )
    if form.get("_return") == "json":
        return {"value": it.id, "label": it.label}
    return redirect(
        "/items", f"Đã thêm {it.code} – {it.name}", request=request
    )


@router.get("/items/export.xlsx")
def items_export(
    request: Request,
    user: CurrentUser = Depends(require("catalog.view")),
    db: Session = Depends(get_tenant_db),
):
    company = db.get(Company, user.company_id)
    book = ReportBook()
    ws = book.add_sheet(
        "Hàng hóa",
        company=company,
        title="DANH MỤC VẬT TƯ, HÀNG HÓA, DỊCH VỤ",
        columns=[
            Col(c, 14 if i not in (1,) else 40)
            for i, c in enumerate(IMPORT_COLUMNS)
        ],
    )
    for it in service.search(
        db, user.company_id, "", active=None, limit=100_000
    ):
        book.row(
            ws,
            [
                it.code,
                it.name,
                it.unit,
                ITEM_KINDS.get(it.kind, it.kind),
                it.vat_rate,
                float(it.sale_price or 0),
                float(it.purchase_price or 0),
                it.stock_account or "",
                it.revenue_account or "",
                it.expense_account or "",
            ],
        )
    return xlsx_response(book.to_bytes(), "hang-hoa.xlsx")


@router.get("/items/import")
def items_import_page(
    request: Request, _u: CurrentUser = Depends(require("catalog.create"))
):
    return render(
        request,
        "items/import.html",
        {
            "columns": IMPORT_COLUMNS,
            "result": None,
            "breadcrumbs": [
                ("Hàng hóa, dịch vụ", "/items"),
                ("Nhập từ Excel", None),
            ],
        },
    )


@router.get("/items/import/template.xlsx")
def items_import_template(
    _u: CurrentUser = Depends(require("catalog.create")),
):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Hàng hóa"
    ws.append(IMPORT_COLUMNS)
    ws.append(
        [
            "DV0001",
            "Dịch vụ tư vấn phần mềm",
            "Giờ",
            "Dịch vụ",
            "10",
            500000,
            0,
            "",
            "5113",
            "",
        ]
    )
    ws.append(
        [
            "HH0001",
            "Máy tính xách tay",
            "Chiếc",
            "Hàng hóa",
            "10",
            25000000,
            21000000,
            "156",
            "5111",
            "",
        ]
    )
    return xlsx_response(save_workbook(wb), "mau-nhap-hang-hoa.xlsx")


def import_rows(db: Session, ctx, data: bytes) -> dict:
    """Import the catalog from Excel (row 1 is the header).

    Nothing is written unless the whole file is valid (one
    transaction).
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise BusinessError("Không đọc được file Excel (.xlsx)") from e
    ws = wb.worksheets[0]
    errors, inputs = [], []
    for n, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if n - 1 > MAX_IMPORT_ROWS:
            limit = f"{MAX_IMPORT_ROWS:,}".replace(",", ".")
            raise BusinessError(
                f"File có quá nhiều dòng (tối đa {limit} dòng mỗi lần nhập)"
            )
        vals = [
            ("" if v is None else str(v).strip()[:500])
            for v in (list(row) + [None] * 10)[:10]
        ]
        if not any(vals):
            continue
        (
            code,
            name,
            unit,
            kind_label,
            rate,
            sale,
            buy,
            stock,
            revenue,
            expense,
        ) = vals
        kind = KIND_BY_LABEL.get(
            kind_label.lower(), "goods" if not kind_label else None
        )
        if kind is None:
            errors.append(f"Dòng {n}: loại '{kind_label}' không hợp lệ")
            continue
        rate = (rate or "10").upper().rstrip("%")
        try:
            if rate.startswith("0.") and rate not in ("0",):
                rate = str(int(Decimal(rate) * 100))
            inputs.append(
                (
                    n,
                    ItemInput(
                        name=name,
                        code=code,
                        unit=unit,
                        kind=kind,
                        vat_rate=rate,
                        sale_price=parse_decimal(sale) or Decimal(0),
                        purchase_price=parse_decimal(buy) or Decimal(0),
                        stock_account=stock or None,
                        revenue_account=revenue or None,
                        expense_account=expense or None,
                    ),
                )
            )
        except (ValueError, ArithmeticError) as e:
            errors.append(f"Dòng {n}: {e}")
    # Assign auto-generated codes once for the whole file (avoids
    # rescanning the entire catalog for every row – O(n²)).
    counters: dict[str, int] = {}
    for _n, inp in inputs:
        if not inp.code.strip() and inp.kind in service.KIND_PREFIX:
            prefix = service.KIND_PREFIX[inp.kind]
            if prefix not in counters:
                counters[prefix] = (
                    int(
                        service.next_code(db, ctx.company_id, inp.kind)[
                            len(prefix) :
                        ]
                    )
                    - 1
                )
            counters[prefix] += 1
            inp.code = f"{prefix}{counters[prefix]:04d}"
    created = 0
    for n, inp in inputs:
        try:
            service.create(db, ctx, inp, commit=False)
            created += 1
        except BusinessError as e:
            errors.append(f"Dòng {n}: {e.message}")
    if errors:
        db.rollback()
        return {"created": 0, "errors": errors}
    db.commit()
    return {"created": created, "errors": []}


@router.post("/items/import")
async def items_import(
    request: Request,
    _u: CurrentUser = Depends(require("catalog.create")),
    db: Session = Depends(get_tenant_db),
):
    form = await request.form()
    f = form.get("file")
    if f is None or not getattr(f, "filename", ""):
        return redirect(
            "/items/import", "Chọn file Excel", "error", request=request
        )
    result = import_rows(db, get_ctx(request), await f.read())
    if result["created"]:
        return redirect(
            "/items",
            f"Đã nhập {result['created']} vật tư, hàng hóa, dịch vụ",
            request=request,
        )
    return render(
        request,
        "items/import.html",
        {
            "columns": IMPORT_COLUMNS,
            "result": result,
            "breadcrumbs": [
                ("Hàng hóa, dịch vụ", "/items"),
                ("Nhập từ Excel", None),
            ],
        },
        status_code=422,
    )


@router.get("/items/{item_id:int}")
def item_detail(
    request: Request,
    item_id: int,
    user: CurrentUser = Depends(require("catalog.view")),
    db: Session = Depends(get_tenant_db),
):
    it = service.get_item(db, user.company_id, item_id)
    return render(
        request,
        "items/form.html",
        _ctx(
            {
                "it": it,
                "form": it,
                "errors": {},
                "breadcrumbs": [
                    ("Hàng hóa, dịch vụ", "/items"),
                    (it.name, None),
                ],
            }
        ),
    )


@router.post("/items/{item_id:int}")
def item_update(
    request: Request,
    item_id: int,
    form: dict = Depends(form_data),
    user: CurrentUser = Depends(require("catalog.manage")),
    db: Session = Depends(get_tenant_db),
):
    it = service.get_item(db, user.company_id, item_id)
    try:
        service.update(db, get_ctx(request), it, _input(form))
    except BusinessError as e:
        db.rollback()
        return redirect(
            f"/items/{item_id}", e.message, "error", request=request
        )
    return redirect(f"/items/{item_id}", "Đã lưu", request=request)
