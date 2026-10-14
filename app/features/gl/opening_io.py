"""Opening balances: grid data, form parsing, Excel template and Excel
import (preview, then confirm).
"""

from __future__ import annotations

import io
import json
import secrets
from dataclasses import asdict, dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.errors import BusinessError
from app.core.money import MoneyError, parse_vnd
from app.core.xlsx import save_workbook
from app.features.accounts.models import Account
from app.features.bank.models import BankAccount
from app.features.journal.opening import OpeningIn, list_opening
from app.features.partners.models import Partner

TEMPLATE_HEADERS = (
    "Số hiệu TK",
    "Mã đối tượng",
    "Số tài khoản ngân hàng",
    "Dư Nợ",
    "Dư Có",
    "Ghi chú",
)


def grid_lines(
    db: Session, company_id: int
) -> tuple[list[dict], dict[int, int]]:
    """Return the grid rows and the balance of each bank account.

    Grid rows cover accounts other than bank accounts; bank balances
    are returned separately, keyed by bank account id.
    """
    rows = list_opening(db, company_id)
    accs = dict(
        db.execute(
            select(Account.code, Account.name).where(
                Account.company_id == company_id
            )
        ).all()
    )
    pids = {r.partner_id for r in rows if r.partner_id}
    partners = (
        {
            p.id: p.label
            for p in db.scalars(
                select(Partner).where(
                    Partner.company_id == company_id, Partner.id.in_(pids)
                )
            )
        }
        if pids
        else {}
    )
    lines, banks = [], {}
    for r in rows:
        if r.bank_account_id:
            banks[r.bank_account_id] = (
                banks.get(r.bank_account_id, 0) + r.debit - r.credit
            )
            continue
        lines.append(
            {
                "account": r.account,
                "account_label": f"{r.account} – {accs.get(r.account, '')}",
                "debit": r.debit or None,
                "credit": r.credit or None,
                "partner_id": r.partner_id or "",
                "partner_label": partners.get(r.partner_id, ""),
                "description": r.note,
            }
        )
    return lines, banks


def parse_form(db: Session, company_id: int, form) -> list[OpeningIn]:
    rows: list[OpeningIn] = []
    errors = []
    try:
        count = min(int(form.get("lines-count") or 0), 2000)
    except ValueError:
        count = 0
    for i in range(count):
        g = lambda f: (form.get(f"lines-{i}-{f}") or "").strip()  # noqa: E731
        acc = g("account")
        if not acc and not g("debit") and not g("credit"):
            continue
        try:
            d = parse_vnd(g("debit")) if g("debit") else 0
            c = parse_vnd(g("credit")) if g("credit") else 0
        except MoneyError:
            errors.append(f"Dòng {i + 1}: số tiền không hợp lệ")
            continue
        pid = g("partner_id")
        rows.append(
            OpeningIn(
                account=acc,
                debit=d,
                credit=c,
                partner_id=int(pid) if pid.isdigit() else None,
                note=g("description"),
            )
        )
    banks = db.scalars(
        select(BankAccount).where(BankAccount.company_id == company_id)
    )
    for b in banks:
        raw = (form.get(f"bank-{b.id}") or "").strip()
        if not raw:
            continue
        try:
            v = parse_vnd(raw)
        except MoneyError:
            errors.append(f"Số dư {b.label} không hợp lệ")
            continue
        if v:
            rows.append(
                OpeningIn(
                    account=b.account_code,
                    debit=max(v, 0),
                    credit=max(-v, 0),
                    bank_account_id=b.id,
                )
            )
    if errors:
        raise BusinessError("; ".join(errors))
    return rows


def template_xlsx(db: Session, company_id: int) -> bytes:
    """Build the Excel template.

    Existing balances (if any) are pre-filled so they can be edited
    and imported again.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "So du dau ky"
    ws.append(list(TEMPLATE_HEADERS))
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEF", (12, 16, 22, 18, 18, 30), strict=True):
        ws.column_dimensions[col].width = w
    rows = list_opening(db, company_id)
    partners = {
        p.id: p.code
        for p in db.scalars(
            select(Partner).where(Partner.company_id == company_id)
        )
    }
    banks = {
        b.id: b.number
        for b in db.scalars(
            select(BankAccount).where(BankAccount.company_id == company_id)
        )
    }
    for r in rows:
        ws.append(
            [
                r.account,
                partners.get(r.partner_id, ""),
                banks.get(r.bank_account_id, ""),
                r.debit or None,
                r.credit or None,
                r.note,
            ]
        )
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.number_format = "#,##0"
    help_ws = wb.create_sheet("Huong dan")
    for line in (
        "Mỗi dòng một số dư. "
        "Cột Dư Nợ / Dư Có: chỉ nhập một bên, số nguyên đồng.",
        "TK theo dõi đối tượng (131, 331, 141, 334…): "
        "nhập Mã đối tượng (theo danh mục Đối tượng).",
        "TK tiền gửi (1121, 1122): "
        "nhập Số tài khoản ngân hàng (theo danh mục Tài khoản ngân hàng).",
        "Nhập file sẽ THAY TOÀN BỘ số dư đầu kỳ hiện có.",
    ):
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 100
    return save_workbook(wb)


@dataclass
class PreviewRow:
    row: int
    account: str
    partner: str
    bank: str
    debit: int
    credit: int
    note: str
    error: str = ""
    partner_id: int | None = None
    bank_account_id: int | None = None


def _num(v) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        raise MoneyError("Số tiền không hợp lệ")
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if v != int(v):
            raise MoneyError("Số tiền phải là số nguyên đồng")
        return int(v)
    return parse_vnd(str(v))


def parse_xlsx(db: Session, company_id: int, data: bytes) -> list[PreviewRow]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # corrupt file or not an .xlsx
        raise BusinessError("Không đọc được file Excel (.xlsx)") from e
    ws = wb.worksheets[0]
    accs = {
        a.code: a
        for a in db.scalars(
            select(Account).where(Account.company_id == company_id)
        )
    }
    partners = {
        p.code.strip().lower(): p
        for p in db.scalars(
            select(Partner).where(Partner.company_id == company_id)
        )
    }
    banks = {
        b.number: b
        for b in db.scalars(
            select(BankAccount).where(BankAccount.company_id == company_id)
        )
    }
    out: list[PreviewRow] = []
    header_seen = False
    for idx, row in enumerate(
        ws.iter_rows(values_only=True, max_col=10), start=1
    ):
        cells = list(row) + [None] * 6
        first = str(cells[0] or "").strip()
        if not header_seen:
            if first.lower().startswith("số hiệu"):
                header_seen = True
            continue
        if not any(c not in (None, "") for c in cells[:6]):
            continue
        if idx > 5000:
            raise BusinessError("File quá nhiều dòng (tối đa 5.000)")
        acc = (
            first.split(".")[0] if first.replace(".", "").isdigit() else first
        )
        pcode = str(cells[1] or "").strip()
        bnum = "".join(str(cells[2] or "").split())
        pr = PreviewRow(
            idx, acc, pcode, bnum, 0, 0, str(cells[5] or "").strip()
        )
        try:
            pr.debit, pr.credit = _num(cells[3]), _num(cells[4])
        except MoneyError as e:
            pr.error = str(e)
        a = accs.get(acc)
        if not pr.error:
            if a is None:
                pr.error = f"Không có TK {acc}"
            elif not a.is_postable:
                pr.error = f"TK {acc} là tài khoản tổng hợp"
            elif pr.debit < 0 or pr.credit < 0 or (pr.debit and pr.credit):
                pr.error = "Chỉ nhập một bên Dư Nợ hoặc Dư Có, số không âm"
        if not pr.error and pcode:
            p = partners.get(pcode.lower())
            if p is None:
                pr.error = f"Không có đối tượng mã {pcode}"
            else:
                pr.partner_id = p.id
                pr.partner = p.label
        if (
            not pr.error
            and a is not None
            and a.partner_mode == "required"
            and not pr.partner_id
        ):
            pr.error = f"TK {acc} cần mã đối tượng"
        if not pr.error and a is not None and a.needs_bank:
            b = banks.get(bnum)
            if b is None:
                pr.error = "Cần số tài khoản ngân hàng có trong danh mục"
            else:
                pr.bank_account_id = b.id
                pr.bank = b.label
        out.append(pr)
    if not header_seen:
        raise BusinessError(
            'File không đúng mẫu: không thấy dòng tiêu đề "Số hiệu TK"'
        )
    return out


def _path(company_id: int, token: str):
    return get_settings().imports_dir / f"opening-{company_id}-{token}.json"


def save_preview(company_id: int, rows: list[PreviewRow]) -> str:
    token = secrets.token_hex(12)
    get_settings().ensure_dirs()
    from app.core.tempfiles import cleanup_imports_throttled

    cleanup_imports_throttled()
    _path(company_id, token).write_text(
        json.dumps([asdict(r) for r in rows], ensure_ascii=False)
    )
    return token


def load_preview(company_id: int, token: str) -> list[PreviewRow]:
    if not token.isalnum():
        raise BusinessError("Phiên nhập không hợp lệ")
    p = _path(company_id, token)
    if not p.exists():
        raise BusinessError("Phiên nhập đã hết hạn, hãy tải file lên lại")
    return [PreviewRow(**r) for r in json.loads(p.read_text())]


def discard_preview(company_id: int, token: str) -> None:
    if token.isalnum():
        _path(company_id, token).unlink(missing_ok=True)


def to_opening(rows: list[PreviewRow]) -> list[OpeningIn]:
    return [
        OpeningIn(
            account=r.account,
            debit=r.debit,
            credit=r.credit,
            partner_id=r.partner_id,
            bank_account_id=r.bank_account_id,
            note=r.note,
        )
        for r in rows
    ]
