"""PDF (WeasyPrint) and Excel (openpyxl) exports for the gl books and
reports.
"""

from __future__ import annotations

from datetime import date

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.pdf import render_pdf
from app.core.xlsx import BOLD, Col, ReportBook
from app.features.gl.statements import statement_set


def date_label(d: date | None = None) -> str:
    d = d or date.today()
    return f"Lập, ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"


def pdf(template: str, **ctx) -> bytes:
    ctx.setdefault("date_label", date_label())
    return render_pdf(f"gl/print/{template}.html", ctx)


# ---------------------------------------------------------------- Excel


def xlsx_statement(company, st, subtitle: str) -> bytes:
    book = ReportBook()
    notes = [Col("Thuyết minh", 11)] if st.has_notes else []
    ws = book.add_sheet(
        st.code,
        company=company,
        form_code=st.code,
        title=st.title,
        subtitle=subtitle,
        columns=[
            Col("CHỈ TIÊU", 62),
            Col("Mã số", 9),
            *notes,
            Col(st.current_label, 20, "money"),
            Col(st.previous_label, 20, "money"),
        ],
    )
    first = book.first_data_row[id(ws)]
    row_of: dict[str, int] = {}
    for i, r in enumerate(st.rows):
        heading = not r.line.code[:1].isdigit()
        idx = first + i
        row_of[r.line.code] = idx
        book.row(
            ws,
            [
                ("    " if r.line.level == 2 else "") + r.line.name,
                "" if heading else r.line.code,
                *([r.line.note] if st.has_notes else []),
                None if heading else r.current,
                None if heading else r.previous,
            ],
            bold=r.line.level < 2,
        )
    # Total lines use formulas so users can verify them in Excel.
    for r in st.rows:
        if not r.line.formula:
            continue
        for col in (4, 5) if st.has_notes else (3, 4):
            letter = get_column_letter(col)
            parts = []
            for f in r.line.formula:
                neg = f.startswith("-")
                ref = row_of.get(f.lstrip("-"))
                if ref:
                    parts.append(("-" if neg else "+") + f"{letter}{ref}")
            if parts:
                expr = "".join(parts)
                ws.cell(
                    row_of[r.line.code],
                    col,
                    "=" + (expr[1:] if expr.startswith("+") else expr),
                )
                ws.cell(row_of[r.line.code], col).font = BOLD
    book.signatures(ws, company, left="Người lập biểu")
    return book.to_bytes()


def xlsx_trial_balance(company, tb, subtitle: str) -> bytes:
    book = ReportBook()
    ws = book.add_sheet(
        "F01-DNN",
        company=company,
        form_code="F01-DNN",
        title=statement_set(company).trial_balance[0].upper(),
        subtitle=subtitle,
        columns=[
            Col("Số hiệu TK", 11),
            Col("Tên tài khoản", 40),
            Col("Dư Nợ đầu kỳ", 17, "money"),
            Col("Dư Có đầu kỳ", 17, "money"),
            Col("Phát sinh Nợ", 17, "money"),
            Col("Phát sinh Có", 17, "money"),
            Col("Dư Nợ cuối kỳ", 17, "money"),
            Col("Dư Có cuối kỳ", 17, "money"),
        ],
    )
    first = book.first_data_row[id(ws)]
    level1_rows = []
    for i, r in enumerate(tb.rows):
        book.row(
            ws,
            [
                r.code,
                ("    " if r.level > 1 else "") + r.name,
                r.open_d,
                r.open_c,
                r.per_d,
                r.per_c,
                r.close_d,
                r.close_c,
            ],
            bold=r.level == 1,
        )
        if r.level == 1:
            level1_rows.append(first + i)
    idx = ws.max_row + 1
    values: list = ["", "Tổng cộng"]
    for col in range(3, 9):
        letter = get_column_letter(col)
        values.append(
            "=" + "+".join(f"{letter}{r}" for r in level1_rows)
            if level1_rows
            else 0
        )
    book._write(ws, idx, values, bold=True)
    book.signatures(ws, company, left="Người lập biểu")
    return book.to_bytes()


def xlsx_journal(company, jb, subtitle: str) -> bytes:
    book = ReportBook()
    ws = book.add_sheet(
        "S03a-DNN",
        company=company,
        form_code="S03a-DNN",
        title="SỔ NHẬT KÝ CHUNG",
        subtitle=subtitle,
        columns=[
            Col("Ngày ghi sổ", 12, "date"),
            Col("Số hiệu CT", 15),
            Col("Ngày CT", 12, "date"),
            Col("Diễn giải", 44),
            Col("STT dòng", 8, "int"),
            Col("Số hiệu TK", 10),
            Col("Nợ", 17, "money"),
            Col("Có", 17, "money"),
        ],
    )
    for r in jb.rows:
        book.row(
            ws,
            [
                r.date,
                r.number,
                r.doc_date,
                r.description,
                r.stt,
                r.account,
                r.debit or None,
                r.credit or None,
            ],
        )
    book.total_row(ws, "Cộng", sum_cols=[6, 7], label_col=4)
    book.signatures(ws, company, left="Người ghi sổ")
    return book.to_bytes()


def xlsx_ledger(
    company, lb, subtitle: str, *, form_code: str | None, title: str
) -> bytes:
    """General ledger / detail ledger.

    The balance column is a running-total formula.
    """
    book = ReportBook()
    ws = book.add_sheet(
        form_code or "So chi tiet",
        company=company,
        form_code=form_code,
        title=title,
        subtitle=subtitle,
        extra_lines=[
            f"Tài khoản: {lb.account} – {lb.account_name}"
            + (f" – Đối tượng: {lb.partner_name}" if lb.partner_name else "")
            + (f" – Ngân hàng: {lb.bank_label}" if lb.bank_label else "")
        ],
        columns=[
            Col("Ngày ghi sổ", 12, "date"),
            Col("Số hiệu CT", 15),
            Col("Ngày CT", 12, "date"),
            Col("Diễn giải", 42),
            Col("TK đối ứng", 10),
            Col("Nợ", 17, "money"),
            Col("Có", 17, "money"),
            Col("Số dư (Nợ − Có)", 18, "money"),
        ],
    )
    first = book.first_data_row[id(ws)]
    book.row(
        ws,
        [None, None, None, "Số dư đầu kỳ", None, None, None, lb.opening],
        bold=True,
    )
    for i, r in enumerate(lb.rows, start=1):
        idx = first + i
        book.row(
            ws,
            [
                r.date,
                r.number,
                r.doc_date or r.date,
                r.description,
                r.contra,
                r.debit or None,
                r.credit or None,
                f"=H{idx - 1}+F{idx}-G{idx}",
            ],
        )
    last = ws.max_row
    idx = last + 1
    book._write(
        ws,
        idx,
        [
            None,
            None,
            None,
            "Cộng số phát sinh",
            None,
            f"=SUM(F{first + 1}:F{last})" if last > first else 0,
            f"=SUM(G{first + 1}:G{last})" if last > first else 0,
            None,
        ],
        bold=True,
    )
    book._write(
        ws,
        idx + 1,
        [
            None,
            None,
            None,
            "Số dư cuối kỳ",
            None,
            None,
            None,
            f"=H{first}+F{idx}-G{idx}",
        ],
        bold=True,
    )
    book.signatures(ws, company, left="Người ghi sổ")
    return book.to_bytes()


def xlsx_notes(company, notes, subtitle: str) -> bytes:
    book = ReportBook()
    ws = book.add_sheet(
        "B09-DNN",
        company=company,
        form_code="B09-DNN",
        title="BẢN THUYẾT MINH BÁO CÁO TÀI CHÍNH",
        subtitle=subtitle,
        columns=[
            Col("Chỉ tiêu", 60),
            Col("Cột 1", 20, "money"),
            Col("Cột 2", 20, "money"),
        ],
    )
    book.row(
        ws,
        ["I. Đặc điểm hoạt động, chế độ kế toán áp dụng", None, None],
        bold=True,
    )
    for k, v in notes.general:
        book.row(ws, [f"{k}: {v}", None, None])
    book.row(ws, ["II. Chính sách kế toán áp dụng", None, None], bold=True)
    for k, v in notes.policies:
        book.row(ws, [f"{k}: {v}", None, None])
    for header, tables in (
        (
            "III. Thông tin bổ sung cho Báo cáo tình hình tài chính",
            notes.balance_tables,
        ),
        (
            "IV. Thông tin bổ sung cho Báo cáo kết quả hoạt động kinh doanh",
            notes.income_tables,
        ),
    ):
        book.row(ws, [header, None, None], bold=True)
        for t in tables:
            book.row(
                ws,
                [f"{t.code}. {t.title}", t.columns[0], t.columns[1]],
                bold=True,
            )
            for label, a, b in t.rows:
                book.row(ws, [label, a, b])
    ws.cell(1, 1).font = Font(name="Times New Roman", bold=True, size=11)
    book.signatures(ws, company, left="Người lập biểu")
    return book.to_bytes()


def xlsx_entries(company, entries, names) -> bytes:
    book = ReportBook()
    ws = book.add_sheet(
        "Chung tu",
        company=company,
        title="DANH SÁCH CHỨNG TỪ",
        columns=[
            Col("Ngày hạch toán", 12, "date"),
            Col("Số chứng từ", 15),
            Col("Loại", 16),
            Col("Diễn giải", 40),
            Col("TK Nợ", 9),
            Col("TK Có", 9),
            Col("Số tiền", 17, "money"),
            Col("Đối tượng", 28),
            Col("Trạng thái", 12),
        ],
    )
    status = {"draft": "Nháp", "posted": "Đã ghi sổ", "cancelled": "Đã hủy"}
    for e in entries:
        for ln in e.lines:
            book.row(
                ws,
                [
                    e.date,
                    e.display_number,
                    e.doc_type_label,
                    ln.description or e.description,
                    ln.debit_account,
                    ln.credit_account,
                    ln.amount,
                    names.partner(ln.debit_partner_id)
                    or names.partner(ln.credit_partner_id),
                    status.get(e.status, e.status),
                ],
            )
    return book.to_bytes()
