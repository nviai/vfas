"""Excel export of the money books S04a-DNN / S04b-DNN / S05-DNN.

Follows the official forms; the balance column is a FORMULA.

The S04a/S04b layout matches app/features/cash/layout.py (shared with
the import): opening balance on row 12, data from row 13.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from app.core.charts.tt133 import ISSUED
from app.core.dates import fmt_date
from app.core.xlsx import save_workbook
from app.features.cash.books import MoneyBook
from app.features.cash.layout import (
    FIRST_DATA_ROW,
    OPENING_ROW,
    S04A,
    S04A_SHEET,
    S04B,
    S04B_SHEET,
    S05,
    S05_SHEET,
)

FONT_NAME = "Times New Roman"
MONEY_FMT = "#,##0"
DATE_FMT = "dd/mm/yyyy"
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# two-level column headers on rows 9-10, column letters on row 11
HEADER_TOP = 9
LETTER_ROW = 11

WIDTHS = {
    "post_date": 12,
    "doc_date": 12,
    "receipt_no": 13,
    "payment_no": 13,
    "number": 13,
    "description": 42,
    "contra": 9,
    "receipt": 15,
    "payment": 15,
    "balance": 16,
    "note": 20,
}


def _font(size: float = 11, bold: bool = False, italic: bool = False) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, italic=italic)


def _put(
    ws,
    row: int,
    col: int,
    value=None,
    *,
    bold=False,
    italic=False,
    size=11,
    align=None,
    wrap=False,
    fmt=None,
    border=False,
    merge_to: int | None = None,
    valign="center",
):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = _font(size, bold, italic)
    cell.alignment = Alignment(
        horizontal=align, vertical=valign, wrap_text=wrap
    )
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = BORDER
    if merge_to and merge_to > col:
        ws.merge_cells(
            start_row=row, start_column=col, end_row=row, end_column=merge_to
        )
    return cell


def book_xlsx(
    company,
    book: MoneyBook,
    period_label: str,
    *,
    subtitle: str = "",
    extra_lines: list[str] | None = None,
) -> bytes:
    form = book.form
    cols = {"s04a": S04A, "s04b": S04B, "s05": S05}[form]
    sheet = {"s04a": S04A_SHEET, "s04b": S04B_SHEET, "s05": S05_SHEET}[form]
    last_col = max(cols.values())
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for key, idx in cols.items():
        ws.column_dimensions[get_column_letter(idx)].width = WIDTHS[key]

    right_start = last_col - 2
    _put(
        ws,
        1,
        1,
        f"Đơn vị: {company.name}",
        bold=True,
        merge_to=right_start - 1,
        align="left",
    )
    _put(
        ws,
        2,
        1,
        f"Địa chỉ: {company.address or ''}",
        merge_to=right_start - 1,
        align="left",
        wrap=True,
    )
    ws.row_dimensions[2].height = 30
    if company.tax_code:
        _put(
            ws,
            3,
            1,
            f"Mã số thuế: {company.tax_code}",
            merge_to=right_start - 1,
            align="left",
        )
    _put(
        ws,
        1,
        right_start,
        f"Mẫu số {book.code}",
        bold=True,
        align="center",
        merge_to=last_col,
    )
    issued = ws.cell(row=2, column=right_start, value=ISSUED)
    issued.font = _font(9, italic=True)
    issued.alignment = Alignment(
        horizontal="center", vertical="top", wrap_text=True
    )
    ws.merge_cells(
        start_row=2, start_column=right_start, end_row=3, end_column=last_col
    )

    _put(
        ws,
        5,
        1,
        book.title,
        bold=True,
        size=14,
        align="center",
        merge_to=last_col,
    )
    ws.row_dimensions[5].height = 22
    _put(ws, 6, 1, subtitle, align="center", merge_to=last_col)
    _put(
        ws, 7, 1, period_label, italic=True, align="center", merge_to=last_col
    )
    _put(
        ws,
        8,
        1,
        "Đơn vị tính: Đồng",
        italic=True,
        size=10,
        align="right",
        merge_to=last_col,
    )
    for i, line in enumerate(extra_lines or []):
        # S05: bank branch and account number, on rows 3 and 4 at
        # the left
        _put(ws, 3 + i, 1, line, merge_to=right_start - 1, align="left")

    top, sub = HEADER_TOP, HEADER_TOP + 1

    def head(col: int, text: str, *, span_rows=False, span_to=None, row=top):
        _put(
            ws,
            row,
            col,
            text,
            bold=True,
            size=10,
            align="center",
            wrap=True,
            border=True,
        )
        if span_rows:
            ws.merge_cells(
                start_row=top, start_column=col, end_row=sub, end_column=col
            )
        elif span_to:
            ws.merge_cells(
                start_row=row,
                start_column=col,
                end_row=row,
                end_column=span_to,
            )

    head(cols["post_date"], "Ngày, tháng ghi sổ", span_rows=True)
    if form == "s05":
        head(cols["number"], "Chứng từ", span_to=cols["doc_date"])
        head(cols["number"], "Số hiệu", row=sub)
        head(cols["doc_date"], "Ngày, tháng", row=sub)
        head(cols["description"], "Diễn giải", span_rows=True)
        head(cols["contra"], "Tài khoản đối ứng", span_rows=True)
        head(cols["receipt"], "Số tiền", span_to=cols["balance"])
        head(cols["receipt"], "Thu (gửi vào)", row=sub)
        head(cols["payment"], "Chi (rút ra)", row=sub)
        head(cols["balance"], "Còn lại", row=sub)
        letters = ["A", "B", "C", "D", "E", "1", "2", "3", "F"]
    else:
        head(cols["doc_date"], "Ngày, tháng chứng từ", span_rows=True)
        head(
            cols["receipt_no"], "Số hiệu chứng từ", span_to=cols["payment_no"]
        )
        head(cols["receipt_no"], "Thu", row=sub)
        head(cols["payment_no"], "Chi", row=sub)
        head(cols["description"], "Diễn giải", span_rows=True)
        if form == "s04b":
            head(cols["contra"], "TK đối ứng", span_rows=True)
            head(cols["receipt"], "Số phát sinh", span_to=cols["payment"])
            head(cols["receipt"], "Nợ", row=sub)
            head(cols["payment"], "Có", row=sub)
            head(cols["balance"], "Số tồn", span_rows=True)
            letters = ["A", "B", "C", "D", "E", "F", "1", "2", "3", "G"]
        else:
            head(cols["receipt"], "Số tiền", span_to=cols["balance"])
            head(cols["receipt"], "Thu", row=sub)
            head(cols["payment"], "Chi", row=sub)
            head(cols["balance"], "Tồn", row=sub)
            letters = ["A", "B", "C", "D", "E", "1", "2", "3", "G"]
    head(cols["note"], "Ghi chú", span_rows=True)
    for r in (top, sub):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).border = BORDER
    ws.row_dimensions[top].height = 30
    for c, letter in enumerate(letters, start=1):
        _put(
            ws,
            LETTER_ROW,
            c,
            letter,
            italic=True,
            size=10,
            align="center",
            border=True,
        )

    in_col = get_column_letter(cols["receipt"])
    out_col = get_column_letter(cols["payment"])
    bal_col = get_column_letter(cols["balance"])

    def body_row(r: int, *, bold=False):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if cell.font is None or cell.font.name != FONT_NAME or bold:
                cell.font = _font(11, bold)

    label = "Số dư đầu kỳ" if form == "s05" else "Số tồn đầu kỳ"
    _put(ws, OPENING_ROW, cols["description"], label, bold=True, italic=True)
    _put(
        ws,
        OPENING_ROW,
        cols["balance"],
        book.opening,
        bold=True,
        fmt=MONEY_FMT,
        align="right",
    )
    body_row(OPENING_ROW)

    r = FIRST_DATA_ROW
    for row in book.rows:
        _put(
            ws,
            r,
            cols["post_date"],
            row.post_date,
            fmt=DATE_FMT,
            align="center",
            valign="top",
        )
        _put(
            ws,
            r,
            cols["doc_date"],
            row.doc_date or row.post_date,
            fmt=DATE_FMT,
            align="center",
            valign="top",
        )
        if form == "s05":
            _put(
                ws,
                r,
                cols["number"],
                row.number or None,
                align="center",
                valign="top",
            )
        else:
            _put(
                ws,
                r,
                cols["receipt_no"],
                row.receipt_no or None,
                align="center",
                valign="top",
            )
            _put(
                ws,
                r,
                cols["payment_no"],
                row.payment_no or None,
                align="center",
                valign="top",
            )
        _put(
            ws,
            r,
            cols["description"],
            row.description,
            wrap=True,
            valign="top",
        )
        if "contra" in cols:
            _put(
                ws,
                r,
                cols["contra"],
                row.contra or None,
                align="center",
                valign="top",
            )
        _put(
            ws,
            r,
            cols["receipt"],
            row.receipt or None,
            fmt=MONEY_FMT,
            valign="top",
        )
        _put(
            ws,
            r,
            cols["payment"],
            row.payment or None,
            fmt=MONEY_FMT,
            valign="top",
        )
        _put(
            ws,
            r,
            cols["balance"],
            f"={bal_col}{r - 1}+{in_col}{r}-{out_col}{r}",
            fmt=MONEY_FMT,
            valign="top",
        )
        _put(
            ws,
            r,
            cols["note"],
            row.note or None,
            size=10,
            wrap=True,
            valign="top",
        )
        body_row(r)
        r += 1
    last_data = r - 1

    total_row = r
    _put(
        ws,
        total_row,
        cols["description"],
        "Cộng phát sinh trong kỳ",
        bold=True,
    )
    if last_data >= FIRST_DATA_ROW:
        in_total = f"=SUM({in_col}{FIRST_DATA_ROW}:{in_col}{last_data})"
        out_total = f"=SUM({out_col}{FIRST_DATA_ROW}:{out_col}{last_data})"
    else:
        in_total, out_total = 0, 0
    _put(ws, total_row, cols["receipt"], in_total, bold=True, fmt=MONEY_FMT)
    _put(ws, total_row, cols["payment"], out_total, bold=True, fmt=MONEY_FMT)
    body_row(total_row, bold=True)
    closing_row = total_row + 1
    _put(
        ws,
        closing_row,
        cols["description"],
        "Số dư cuối kỳ" if form == "s05" else "Số tồn cuối kỳ",
        bold=True,
    )
    _put(
        ws,
        closing_row,
        cols["balance"],
        f"={bal_col}{OPENING_ROW}+{in_col}{total_row}-{out_col}{total_row}",
        bold=True,
        fmt=MONEY_FMT,
    )
    body_row(closing_row, bold=True)

    f = closing_row + 2
    _put(
        ws,
        f,
        1,
        "- Sổ này có ..... trang, đánh số từ trang 01 đến trang .....",
        merge_to=cols["description"],
        align="left",
    )
    _put(
        ws,
        f + 1,
        1,
        f"- Ngày mở sổ: {fmt_date(company.start_date)}",
        merge_to=cols["description"],
        align="left",
    )
    sign_row = f + 3
    third = last_col // 3
    blocks = [
        (1, third, "Người ghi sổ", "(Ký, họ tên)", company.preparer),
        (
            third + 1,
            2 * third,
            "Kế toán trưởng",
            "(Ký, họ tên)",
            company.chief_accountant,
        ),
        (
            2 * third + 1,
            last_col,
            "Người đại diện theo pháp luật",
            "(Ký, họ tên, đóng dấu)",
            company.legal_rep,
        ),
    ]
    _put(
        ws,
        sign_row,
        blocks[2][0],
        "Ngày ..... tháng ..... năm ..........",
        italic=True,
        align="center",
        merge_to=last_col,
    )
    for start, end, role, hint, name in blocks:
        _put(
            ws,
            sign_row + 1,
            start,
            role,
            bold=True,
            align="center",
            merge_to=end,
        )
        _put(
            ws,
            sign_row + 2,
            start,
            hint,
            italic=True,
            size=10,
            align="center",
            merge_to=end,
        )
        _put(
            ws,
            sign_row + 7,
            start,
            name or None,
            bold=True,
            align="center",
            merge_to=end,
        )

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.3, top=0.5, bottom=0.6)
    ws.print_title_rows = f"{HEADER_TOP}:{LETTER_ROW}"
    ws.oddFooter.right.text = "Trang &P/&N"
    ws.oddFooter.right.size = 8
    ws.print_options.horizontalCentered = True

    return save_workbook(wb)


def template_xlsx(company) -> bytes:
    """Build the cash book import template (S04a-DNN + S04b-DNN).

    Rows noted "Dòng ví dụ" are skipped on import.
    """
    from openpyxl import load_workbook

    from app.features.cash.books import BookRow
    from app.features.cash.layout import EXAMPLE_ROW_MARKER

    d = company.start_date
    example = BookRow(
        post_date=d,
        doc_date=d,
        receipt_no="PT0000",
        payment_no="",
        description="Thu tiền bán hàng (ví dụ cách ghi)",
        contra="5111",
        receipt=1_000_000,
        payment=0,
        balance=1_000_000,
        note=EXAMPLE_ROW_MARKER,
    )
    book_a = MoneyBook(
        form="s04a", account="1111", start=d, end=d, opening=0, rows=[example]
    )
    book_b = MoneyBook(
        form="s04b", account="1111", start=d, end=d, opening=0, rows=[example]
    )
    wb = load_workbook(
        BytesIO(
            book_xlsx(
                company,
                book_a,
                "Từ ngày … đến ngày …",
                subtitle="Loại quỹ: Tiền Việt Nam",
            )
        )
    )
    wb_b = load_workbook(
        BytesIO(
            book_xlsx(
                company,
                book_b,
                "Từ ngày … đến ngày …",
                subtitle="Tài khoản: 1111 – Loại quỹ: Tiền Việt Nam",
            )
        )
    )
    src = wb_b[S04B_SHEET]
    dst = wb.create_sheet(S04B_SHEET)
    for row in src.iter_rows():
        for cell in row:
            new = dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new.font = cell.font.copy()
                new.border = cell.border.copy()
                new.alignment = cell.alignment.copy()
                new.number_format = cell.number_format
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
    return save_workbook(wb)
