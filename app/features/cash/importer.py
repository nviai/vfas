"""Import an existing cash book from Excel as posted vouchers.

Reads S04a-DNN (plus S04b-DNN for the offsetting accounts) and creates
posted cash receipts/payments.

Flow: upload the file -> read and validate (preview) -> the user
confirms -> everything is written in ONE transaction (the company is
locked and everything is validated again; the preview result is not
trusted). The original voucher numbers are kept.

Historical data often lacks details that double-entry bookkeeping
needs:
- missing offsetting account (no S04b): the user picks a default
  account for receipt rows and for payment rows;
- accounts that require a partner (131, 331, 141...): the placeholder
  partner `UNKNOWN_PARTNER` is attached, to be fixed later;
- bank accounts (1121): the bank account chosen by the user is
  attached (a temporary bank account is created if there is none).
The retroactive negative-balance check covers TK 1111 only
(historical bank balances are often not in the software yet).
"""

from __future__ import annotations

import json
import math
import re
import secrets
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.audit import write_audit
from app.core.config import get_settings
from app.core.dates import fmt_date, parse_date, today
from app.core.errors import BusinessError, NotFoundError
from app.core.money import MAX_MONEY, MoneyError, format_vnd, parse_vnd
from app.features.accounts.models import Account
from app.features.bank.models import BankAccount
from app.features.cash.layout import (
    EXAMPLE_ROW_MARKER,
    FIRST_DATA_ROW,
    OPENING_ROW,
    S04A,
    S04A_SHEET,
    S04B,
    S04B_SHEET,
)
from app.features.journal import locks, queries
from app.features.journal import service as js
from app.features.journal.models import JournalEntry, OpeningBalance
from app.features.org.models import Company
from app.features.partners.models import Partner

CASH = "1111"
TOKEN_RE = re.compile(r"^[0-9a-f]{32}$")
NUMBER_RE = re.compile(r"^(.*?)(\d+)$")
STOP_MARKERS = ("cộng phát sinh", "số tồn cuối kỳ")
MAX_EMPTY_RUN = 20
TEMP_MAX_AGE = 2 * 86400
UNKNOWN_PARTNER = "Đối tượng chưa xác định (nhập từ sổ quỹ)"
RECEIPT, PAYMENT = "PT", "PC"


@dataclass
class Issue:
    """An error or warning.

    ``msg`` is the Vietnamese source string (the i18n key) and
    ``params`` its format arguments.
    """

    msg: str
    params: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return self.msg.format(**self.params) if self.params else self.msg


@dataclass
class ImportRow:
    row_no: int
    post_date: date | None = None
    doc_date: date | None = None
    number: str | None = None
    kind: str | None = None
    amount: int | None = None
    description: str = ""
    note: str = ""
    contra: str | None = None
    file_balance: int | None = None
    balance: int | None = None
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    @property
    def receipt(self) -> int:
        return (
            self.amount
            if self.ok and self.kind == RECEIPT and self.amount
            else 0
        )

    @property
    def payment(self) -> int:
        return (
            self.amount
            if self.ok and self.kind == PAYMENT and self.amount
            else 0
        )

    @property
    def signed(self) -> int:
        return (
            (self.amount or 0) if self.kind == RECEIPT else -(self.amount or 0)
        )


@dataclass
class ImportOptions:
    default_receipt_contra: str = ""
    default_payment_contra: str = ""
    bank_account_id: int | None = None


@dataclass
class ImportPreview:
    filename: str
    options: ImportOptions = field(default_factory=ImportOptions)
    rows: list[ImportRow] = field(default_factory=list)
    blocking: list[Issue] = field(default_factory=list)
    notices: list[Issue] = field(default_factory=list)
    skipped_examples: int = 0
    has_s04b: bool = False
    opening_file: int | None = None
    # set: set the TK 1111 opening balance to H12;
    # check: compare it with the books
    opening_mode: str = "set"
    opening_current: int = 0
    opening_expected: int | None = None
    file_total_in: int | None = None
    file_total_out: int | None = None
    file_closing: int | None = None
    closing_after: int = 0
    # kind -> number of rows missing an offsetting account
    missing_contra: dict[str, int] = field(default_factory=dict)
    needs_bank: bool = False
    partner_placeholder: int = 0

    @property
    def valid_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if r.ok]

    @property
    def error_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if not r.ok]

    @property
    def warning_rows(self) -> list[ImportRow]:
        return [r for r in self.rows if r.ok and r.warnings]

    @property
    def total_in(self) -> int:
        return sum(r.receipt for r in self.rows)

    @property
    def total_out(self) -> int:
        return sum(r.payment for r in self.rows)

    @property
    def can_import(self) -> bool:
        return not self.blocking and bool(self.valid_rows)


# -------------------------------------------------------- reading cells


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _date(v) -> tuple[date | None, bool]:
    if _is_blank(v):
        return None, True
    if isinstance(v, datetime):
        return v.date(), True
    if isinstance(v, date):
        return v, True
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:
            try:
                return from_excel(v).date(), True
            except (ValueError, OverflowError):
                return None, False
        return None, False
    try:
        return parse_date(str(v)), True
    except ValueError:
        return None, False


def _amount(v) -> tuple[int | None, Issue | None]:
    if _is_blank(v):
        return None, None
    if isinstance(v, bool):
        return None, Issue("Số tiền không hợp lệ")
    if isinstance(v, int):
        value = v
    elif isinstance(v, (float, Decimal)):
        if not _finite(v):
            return None, Issue(
                "Số tiền không hợp lệ: {value}", {"value": str(v)[:30]}
            )
        if v != int(v):
            return None, Issue(
                "Số tiền {value} không phải số nguyên đồng", {"value": str(v)}
            )
        value = int(v)
    else:
        try:
            value = parse_vnd(str(v))
        except MoneyError:
            return None, Issue(
                "Số tiền không hợp lệ: {value}", {"value": str(v).strip()}
            )
    if value < 0:
        return None, Issue("Số tiền âm: {value}", {"value": format_vnd(value)})
    if value > MAX_MONEY:
        return None, Issue("Số tiền quá lớn: {value}", {"value": str(v)[:30]})
    return value, None


def _finite(v) -> bool:
    if isinstance(v, float):
        return math.isfinite(v) and abs(v) <= MAX_MONEY
    if isinstance(v, Decimal):
        return v.is_finite() and abs(v) <= MAX_MONEY
    return True


def _balance(v) -> int | None:
    if _is_blank(v) or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v if abs(v) <= MAX_MONEY else None
    if isinstance(v, (float, Decimal)):
        if not _finite(v):
            return None
        return int(v) if v == int(v) else None
    s = str(v).strip()
    neg = s.startswith("-")
    try:
        value = parse_vnd(s.lstrip("-").strip())
    except MoneyError:
        return None
    return -value if neg else value


def _find_sheet(wb, name: str):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    return None


def _cell(values: tuple, col: int):
    return values[col - 1] if len(values) >= col else None


def _iter_data(ws, ncols: int):
    empty_run = 0
    for row_no, values in enumerate(
        ws.iter_rows(min_row=FIRST_DATA_ROW, max_col=ncols, values_only=True),
        start=FIRST_DATA_ROW,
    ):
        if all(_is_blank(v) for v in values):
            empty_run += 1
            if empty_run >= MAX_EMPTY_RUN:
                return
            continue
        empty_run = 0
        yield row_no, values


def _stop_kind(description: str) -> str | None:
    low = description.strip().lower()
    for marker in STOP_MARKERS:
        if low.startswith(marker):
            return marker
    return None


# ----------------------------------------------------- reading the file


def _read_workbook(path: Path, preview: ImportPreview) -> None:
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - corrupt file or not an xlsx
        preview.blocking.append(
            Issue("Không đọc được file. Hãy chọn file Excel .xlsx.")
        )
        return
    try:
        ws = _find_sheet(wb, S04A_SHEET)
        if ws is None:
            preview.blocking.append(
                Issue(
                    "Không tìm thấy sheet {sheet} trong file",
                    {"sheet": S04A_SHEET},
                )
            )
            return
        opening_values = next(
            ws.iter_rows(
                min_row=OPENING_ROW,
                max_row=OPENING_ROW,
                max_col=len(S04A),
                values_only=True,
            ),
            (),
        )
        raw_opening = _cell(opening_values, S04A["balance"])
        if raw_opening is None:
            raw_opening = _raw_opening(path)
        if raw_opening is None:
            preview.opening_file = 0
            preview.notices.append(
                Issue("Ô H12 (tồn đầu kỳ) trống – coi tồn đầu kỳ của file = 0")
            )
        elif isinstance(raw_opening, str) and raw_opening.startswith("="):
            preview.blocking.append(
                Issue(
                    "Ô H12 là công thức chưa được tính ({formula}). "
                    "Mở file bằng Excel, lưu lại rồi tải lên lần nữa.",
                    {"formula": raw_opening},
                )
            )
        else:
            opening = _balance(raw_opening)
            if opening is None:
                preview.blocking.append(
                    Issue(
                        "Không đọc được tồn đầu kỳ ở ô H12: {value}",
                        {"value": str(raw_opening)},
                    )
                )
            else:
                preview.opening_file = opening

        contra_map = _read_contra(wb)
        preview.has_s04b = contra_map is not None
        if contra_map is None:
            preview.notices.append(
                Issue(
                    "Không có sheet {sheet}: chọn TK đối ứng mặc định "
                    "cho các dòng thu, chi",
                    {"sheet": S04B_SHEET},
                )
            )
            contra_map = {}

        last_balance = None
        for row_no, values in _iter_data(ws, len(S04A)):
            description = _text(_cell(values, S04A["description"]))
            stop = _stop_kind(description)
            if stop == "cộng phát sinh":
                preview.file_total_in = _balance(
                    _cell(values, S04A["receipt"])
                )
                preview.file_total_out = _balance(
                    _cell(values, S04A["payment"])
                )
                continue
            if stop == "số tồn cuối kỳ":
                preview.file_closing = _balance(_cell(values, S04A["balance"]))
                break
            note = _text(_cell(values, S04A["note"]))
            if EXAMPLE_ROW_MARKER.lower() in note.lower():
                preview.skipped_examples += 1
                continue
            row = _parse_row(row_no, values, description, note, contra_map)
            if row.file_balance is not None:
                last_balance = row.file_balance
            preview.rows.append(row)
        if preview.file_closing is None:
            preview.file_closing = last_balance
    finally:
        wb.close()


def _raw_opening(path: Path):
    try:
        wb = load_workbook(path, data_only=False, read_only=True)
    except Exception:  # noqa: BLE001
        return None
    try:
        ws = _find_sheet(wb, S04A_SHEET)
        if ws is None:
            return None
        values = next(
            ws.iter_rows(
                min_row=OPENING_ROW,
                max_row=OPENING_ROW,
                max_col=len(S04A),
                values_only=True,
            ),
            (),
        )
        return _cell(values, S04A["balance"])
    finally:
        wb.close()


def _read_contra(wb) -> dict[str, str] | None:
    ws = _find_sheet(wb, S04B_SHEET)
    if ws is None:
        return None
    result: dict[str, str] = {}
    for _row_no, values in _iter_data(ws, len(S04B)):
        desc = _text(_cell(values, S04B["description"]))
        if _stop_kind(desc):
            if desc.lower().startswith("số tồn cuối kỳ"):
                break
            continue
        if (
            EXAMPLE_ROW_MARKER.lower()
            in _text(_cell(values, S04B["note"])).lower()
        ):
            continue
        contra = _text(_cell(values, S04B["contra"]))
        if not contra:
            continue
        for col in ("receipt_no", "payment_no"):
            number = _text(_cell(values, S04B[col]))
            if number:
                result.setdefault(number, contra)
    return result


def _parse_row(
    row_no: int, values: tuple, description: str, note: str, contra_map: dict
) -> ImportRow:
    row = ImportRow(row_no=row_no, description=description, note=note)
    post_date, ok = _date(_cell(values, S04A["post_date"]))
    if not ok:
        row.errors.append(
            Issue(
                "Ngày ghi sổ không hợp lệ: {value}",
                {"value": _text(_cell(values, S04A["post_date"]))},
            )
        )
    elif post_date is None:
        row.errors.append(Issue("Thiếu ngày ghi sổ"))
    row.post_date = post_date
    doc_date, ok = _date(_cell(values, S04A["doc_date"]))
    if not ok:
        row.warnings.append(
            Issue("Ngày chứng từ không hợp lệ, dùng ngày ghi sổ")
        )
        doc_date = None
    if doc_date is None:
        if ok:
            row.warnings.append(Issue("Thiếu ngày chứng từ, dùng ngày ghi sổ"))
        doc_date = post_date
    elif post_date and doc_date > post_date:
        row.warnings.append(
            Issue("Ngày chứng từ sau ngày ghi sổ – dùng ngày ghi sổ")
        )
        doc_date = post_date
    row.doc_date = doc_date

    receipt, err_in = _amount(_cell(values, S04A["receipt"]))
    payment, err_out = _amount(_cell(values, S04A["payment"]))
    for err in (err_in, err_out):
        if err:
            row.errors.append(err)
    has_in, has_out = bool(receipt), bool(payment)
    if has_in and has_out:
        row.errors.append(
            Issue("Vừa có số thu vừa có số chi trên cùng một dòng")
        )
    elif not has_in and not has_out and not (err_in or err_out):
        row.errors.append(Issue("Không có số tiền thu hoặc chi"))
    if has_in and not has_out:
        row.kind, row.amount = RECEIPT, receipt
    elif has_out and not has_in:
        row.kind, row.amount = PAYMENT, payment

    receipt_no = _text(_cell(values, S04A["receipt_no"]))
    payment_no = _text(_cell(values, S04A["payment_no"]))
    if row.kind == RECEIPT:
        row.number = receipt_no or None
        if not receipt_no:
            row.errors.append(Issue("Thiếu số phiếu thu (cột C)"))
        if payment_no:
            row.warnings.append(
                Issue(
                    "Dòng thu nhưng có số phiếu chi {number} ở cột D – bỏ qua",
                    {"number": payment_no},
                )
            )
    elif row.kind == PAYMENT:
        row.number = payment_no or None
        if not payment_no:
            row.errors.append(Issue("Thiếu số phiếu chi (cột D)"))
        if receipt_no:
            row.warnings.append(
                Issue(
                    "Dòng chi nhưng có số phiếu thu {number} ở cột C – bỏ qua",
                    {"number": receipt_no},
                )
            )
    else:
        row.number = receipt_no or payment_no or None
        if not row.number:
            row.errors.append(Issue("Thiếu số phiếu"))
    if row.number and len(row.number) > 40:
        row.errors.append(Issue("Số phiếu dài quá 40 ký tự"))
    if not description:
        row.warnings.append(Issue("Thiếu diễn giải"))
    row.file_balance = _balance(_cell(values, S04A["balance"]))
    row.contra = contra_map.get(row.number) if row.number else None
    return row


# -------------------------------------- validation against the database


def _detail_child(accounts: dict[str, Account], code: str) -> str | None:
    """Pick a detail account to replace a summary account.

    Prefers an "other" ("khác") child account, otherwise the first
    child account.
    """
    children = sorted(
        (
            a
            for a in accounts.values()
            if a.code.startswith(code)
            and a.code != code
            and a.is_postable
            and a.is_active
        ),
        key=lambda a: a.code,
    )
    if not children:
        return None
    direct = [a for a in children if a.parent_code == code] or children
    other = [a for a in direct if "khác" in a.name.lower()]
    return (other or direct)[0].code


def _cash_posted_exists(db: Session, company_id: int) -> bool:
    lr = queries.ledger_source(db, company_id, include_opening=False)
    return bool(
        db.scalar(
            select(func.count()).select_from(lr).where(lr.c.account == CASH)
        )
    )


def _validate_against_db(
    db: Session, company: Company, preview: ImportPreview
) -> None:
    t = today()
    locked = locks.locked_set(db, company.id)
    existing_numbers = set(
        db.scalars(
            select(JournalEntry.number).where(
                JournalEntry.company_id == company.id,
                JournalEntry.number.is_not(None),
            )
        )
    )
    accounts = {
        a.code: a
        for a in db.scalars(
            select(Account).where(Account.company_id == company.id)
        )
    }
    opts = preview.options
    defaults = {
        RECEIPT: opts.default_receipt_contra.strip(),
        PAYMENT: opts.default_payment_contra.strip(),
    }
    for kind, code in defaults.items():
        acc = accounts.get(code)
        if code and (
            acc is None
            or not acc.is_postable
            or not acc.is_active
            or code == CASH
        ):
            preview.blocking.append(
                Issue(
                    "TK đối ứng mặc định {code} không hợp lệ "
                    "(phải là tài khoản chi tiết)",
                    {"code": code},
                )
            )
            defaults[kind] = ""
    seen: dict[str, int] = {}
    for row in preview.rows:
        d = row.post_date
        if d is not None:
            if d < company.start_date:
                row.errors.append(
                    Issue(
                        "Ngày ghi sổ {date} trước ngày bắt đầu "
                        "hạch toán {start}",
                        {
                            "date": fmt_date(d),
                            "start": fmt_date(company.start_date),
                        },
                    )
                )
            elif d > t:
                row.errors.append(
                    Issue(
                        "Ngày ghi sổ {date} ở tương lai", {"date": fmt_date(d)}
                    )
                )
            elif locks.month_locked(locked, d.year, d.month):
                row.errors.append(
                    Issue(
                        "Kỳ {period} đã khóa sổ",
                        {"period": f"{d.month:02d}/{d.year}"},
                    )
                )
        if row.number:
            if row.number in existing_numbers:
                row.errors.append(
                    Issue(
                        "Số phiếu {number} đã có trong sổ",
                        {"number": row.number},
                    )
                )
            elif row.number in seen:
                row.errors.append(
                    Issue(
                        "Trùng số phiếu với dòng {row} trong file",
                        {"row": seen[row.number]},
                    )
                )
            else:
                seen[row.number] = row.row_no
        if row.contra:
            acc = accounts.get(row.contra)
            if acc is not None and not acc.is_postable:
                child = _detail_child(accounts, row.contra)
                if child:
                    row.warnings.append(
                        Issue(
                            "TK {code} là TK tổng hợp – "
                            "dùng TK chi tiết {child}",
                            {"code": row.contra, "child": child},
                        )
                    )
                    row.contra = child
                    acc = accounts[child]
            if acc is None or not acc.is_postable or row.contra == CASH:
                row.warnings.append(
                    Issue(
                        "TK đối ứng {code} không hợp lệ "
                        "(không có hoặc là TK tổng hợp)",
                        {"code": row.contra},
                    )
                )
                row.contra = None
        if not row.contra and row.kind:
            if defaults.get(row.kind):
                row.contra = defaults[row.kind]
                row.warnings.append(
                    Issue(
                        "Dùng TK đối ứng mặc định {code}", {"code": row.contra}
                    )
                )
            else:
                preview.missing_contra[row.kind] = (
                    preview.missing_contra.get(row.kind, 0) + 1
                )
                row.errors.append(
                    Issue(
                        "Thiếu TK đối ứng – chọn TK mặc định ở trên "
                        "hoặc bổ sung sheet S04b-DNN"
                    )
                )
        if row.ok and row.contra:
            acc = accounts[row.contra]
            if acc.needs_bank:
                preview.needs_bank = True
            if acc.partner_mode == "required":
                preview.partner_placeholder += 1
                row.warnings.append(
                    Issue(
                        'TK {code} cần đối tượng – gắn tạm "{name}"',
                        {"code": row.contra, "name": UNKNOWN_PARTNER},
                    )
                )
    if preview.needs_bank:
        bank = (
            db.get(BankAccount, opts.bank_account_id)
            if opts.bank_account_id
            else None
        )
        if bank is None or bank.company_id != company.id:
            has_any = db.scalar(
                select(func.count(BankAccount.id)).where(
                    BankAccount.company_id == company.id
                )
            )
            if has_any:
                preview.blocking.append(
                    Issue(
                        "Chọn tài khoản ngân hàng cho các dòng đối ứng TK 112"
                    )
                )
            else:
                preview.notices.append(
                    Issue(
                        "Chưa có tài khoản ngân hàng: hệ thống tạo "
                        'tài khoản tạm "Ngân hàng (nhập từ sổ quỹ)" – '
                        "đổi tên sau trong danh mục"
                    )
                )
        preview.notices.append(
            Issue(
                "Dòng đối ứng TK 112: số dư tiền gửi có thể âm "
                "nếu chưa nhập số dư đầu kỳ ngân hàng"
            )
        )

    has_posted = _cash_posted_exists(db, company.id)
    opening_now = int(
        db.scalar(
            select(
                func.coalesce(
                    func.sum(OpeningBalance.debit - OpeningBalance.credit), 0
                )
            ).where(
                OpeningBalance.company_id == company.id,
                OpeningBalance.account == CASH,
            )
        )
        or 0
    )
    preview.opening_current = opening_now
    preview.opening_mode = "check" if has_posted else "set"
    if (
        not has_posted
        and preview.opening_file is not None
        and preview.opening_file != opening_now
        and locks.any_locked(db, company.id)
    ):
        preview.blocking.append(
            Issue(
                "Đã khóa sổ nên không đặt được số dư đầu kỳ TK 1111 theo file"
            )
        )
    if has_posted:
        dates = [r.post_date for r in preview.rows if r.ok and r.post_date]
        if dates and preview.opening_file is not None:
            first = min(dates)
            preview.opening_expected = queries.balance(
                db, company.id, CASH, before=first, exact=True
            )
            if preview.opening_expected != preview.opening_file:
                preview.blocking.append(
                    Issue(
                        "Tồn đầu kỳ trong file (ô H12) là {file} "
                        "nhưng tồn trên sổ trước ngày {date} là {book}. "
                        "Không thể nhập vì số liệu không nối tiếp.",
                        {
                            "file": format_vnd(preview.opening_file),
                            "date": fmt_date(first),
                            "book": format_vnd(preview.opening_expected),
                        },
                    )
                )
    _check_negative(db, company.id, preview)


def _check_negative(
    db: Session, company_id: int, preview: ImportPreview
) -> None:
    """Check that the new rows never make the cash balance negative.

    Merges the existing TK 1111 ledger with the new rows in posting
    order and rejects, one by one, new payments that drive the cash
    balance negative.
    """
    opening = (
        preview.opening_file
        if preview.opening_mode == "set" and preview.opening_file is not None
        else preview.opening_current
    )
    lr = queries.ledger_source(db, company_id, include_opening=False)
    existing = db.execute(
        select(
            lr.c.date,
            lr.c.posting_seq,
            lr.c.line_no,
            (lr.c.debit - lr.c.credit),
        )
        .where(lr.c.account == CASH)
        .order_by(lr.c.date, lr.c.posting_seq, lr.c.line_no)
    ).all()
    entries: list[tuple[tuple, int, ImportRow | None]] = [
        ((d, 0, seq or 0, ln or 0), int(amt), None)
        for d, seq, ln, amt in existing
    ]
    for idx, row in enumerate(preview.rows):
        if row.ok and row.post_date and row.amount:
            entries.append(((row.post_date, 1, idx, 0), row.signed, row))
    entries.sort(key=lambda e: e[0])
    excluded: set[int] = set()
    while True:
        bal = opening
        violation = culprit = None
        for i, (_key, signed, row) in enumerate(entries):
            if row is not None and id(row) in excluded:
                continue
            bal += signed
            if row is not None and row.kind == PAYMENT:
                culprit = row
            if bal < 0:
                violation = (i, bal)
                break
        if violation is None:
            break
        if culprit is None:
            preview.notices.append(
                Issue("Sổ quỹ hiện có đã có dòng tồn âm (không do file này)")
            )
            break
        at_date = entries[violation[0]][0][0]
        culprit.errors.append(
            Issue(
                "Gây âm quỹ: tồn tại ngày {date} sẽ còn {balance}",
                {
                    "date": fmt_date(at_date),
                    "balance": format_vnd(violation[1]),
                },
            )
        )
        excluded.add(id(culprit))
    bal = opening
    for _key, signed, row in entries:
        if row is not None and not row.ok:
            continue
        bal += signed
        if row is not None:
            row.balance = bal
            if row.file_balance is not None and row.file_balance != bal:
                row.warnings.append(
                    Issue(
                        "Cột Tồn trong file là {file}, tính lại được {calc}",
                        {
                            "file": format_vnd(row.file_balance),
                            "calc": format_vnd(bal),
                        },
                    )
                )
    preview.closing_after = bal


def build_preview(
    db: Session,
    company: Company,
    path: Path,
    filename: str,
    options: ImportOptions | None = None,
) -> ImportPreview:
    preview = ImportPreview(
        filename=filename, options=options or ImportOptions()
    )
    _read_workbook(path, preview)
    if preview.blocking and not preview.rows:
        return preview
    _validate_against_db(db, company, preview)
    return preview


# ---------------------------------------------- writing to the database


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _series_seq(number: str) -> int | None:
    m = NUMBER_RE.match(number)
    return int(m.group(2)) if m else None


def _placeholder_partner(db: Session, ctx) -> Partner:
    from app.features.partners.service import PartnerInput, create

    p = db.scalar(
        select(Partner).where(
            Partner.company_id == ctx.company_id,
            Partner.name == UNKNOWN_PARTNER,
        )
    )
    return p or create(
        db,
        ctx,
        PartnerInput(
            name=UNKNOWN_PARTNER, note="Tạo khi nhập sổ quỹ từ Excel"
        ),
        commit=False,
    )


def _placeholder_bank(db: Session, ctx) -> BankAccount:
    b = db.scalar(
        select(BankAccount).where(
            BankAccount.company_id == ctx.company_id,
            BankAccount.number == "CHUA-XAC-DINH",
        )
    )
    if b is None:
        b = BankAccount(
            company_id=ctx.company_id,
            account_code="1121",
            number="CHUA-XAC-DINH",
            bank_name="Ngân hàng (nhập từ sổ quỹ)",
            short_name="NH tạm",
            is_active=True,
        )
        db.add(b)
        db.flush()
    return b


def confirm_import(
    db: Session,
    ctx,
    path: Path,
    filename: str,
    *,
    options: ImportOptions | None = None,
    skip_errors: bool,
    token: str | None = None,
) -> ImportPreview:
    """Write all valid rows in one transaction (company locked).

    Return the preview that was used for writing.
    """
    company = db.get(Company, ctx.company_id)
    with js.company_tx(db, ctx.company_id):
        preview = build_preview(db, company, path, filename, options)
        if preview.blocking:
            raise BusinessError(
                "Không thể nhập: "
                + "; ".join(i.text for i in preview.blocking)
            )
        if not preview.valid_rows:
            raise BusinessError("Không có dòng hợp lệ để nhập")
        if preview.error_rows and not skip_errors:
            raise BusinessError(
                f"File còn {len(preview.error_rows)} dòng lỗi. "
                "Đánh dấu ô xác nhận bỏ qua dòng lỗi nếu muốn "
                "chỉ nhập các dòng hợp lệ."
            )
        opening_set = False
        if (
            preview.opening_mode == "set"
            and preview.opening_file is not None
            and preview.opening_file != preview.opening_current
        ):
            db.execute(
                delete(OpeningBalance).where(
                    OpeningBalance.company_id == ctx.company_id,
                    OpeningBalance.account == CASH,
                )
            )
            if preview.opening_file:
                db.add(
                    OpeningBalance(
                        company_id=ctx.company_id,
                        account=CASH,
                        debit=max(preview.opening_file, 0),
                        credit=max(-preview.opening_file, 0),
                        note=f"Nhập từ Excel: {filename}",
                    )
                )
            opening_set = True
            write_audit(
                db,
                ctx,
                action="opening.update",
                entity_type="opening",
                entity_id=CASH,
                summary="Đặt số dư đầu kỳ TK 1111 theo file nhập sổ quỹ",
                before={"1111": preview.opening_current},
                after={"1111": preview.opening_file},
                reason=f"Nhập từ Excel: {filename}",
            )

        accounts = {
            a.code: a
            for a in db.scalars(
                select(Account).where(Account.company_id == ctx.company_id)
            )
        }
        partner = bank = None
        created: list[JournalEntry] = []
        # Post in date order (keeping the file order within a day)
        ordered = sorted(
            ((r.post_date, i, r) for i, r in enumerate(preview.valid_rows)),
            key=lambda x: (x[0], x[1]),
        )
        for _d, _i, row in ordered:
            acc = accounts[row.contra]
            pid = bid = None
            if acc.partner_mode == "required":
                partner = partner or _placeholder_partner(db, ctx)
                pid = partner.id
            if acc.needs_bank:
                if bank is None:
                    opt = preview.options.bank_account_id
                    b = db.get(BankAccount, opt) if opt else None
                    bank = (
                        b
                        if b is not None and b.company_id == ctx.company_id
                        else _placeholder_bank(db, ctx)
                    )
                bid = bank.id
            if row.kind == RECEIPT:
                line = js.LineIn(
                    debit_account=CASH,
                    credit_account=row.contra,
                    amount=row.amount,
                    description=row.description,
                    credit_partner_id=pid,
                    credit_bank_account_id=bid,
                )
            else:
                line = js.LineIn(
                    debit_account=row.contra,
                    credit_account=CASH,
                    amount=row.amount,
                    description=row.description,
                    debit_partner_id=pid,
                    debit_bank_account_id=bid,
                )
            meta = {"imported": filename, "row": row.row_no}
            if row.note:
                meta["note"] = row.note
            e = js.create_entry(
                db,
                ctx,
                js.EntryIn(
                    doc_type=row.kind,
                    date=row.post_date,
                    doc_date=row.doc_date or row.post_date,
                    lines=[line],
                    description=row.description
                    or ("Thu tiền" if row.kind == RECEIPT else "Chi tiền"),
                    number=row.number,
                    source="import",
                    meta=meta,
                ),
                commit=False,
            )
            e.number_seq = _series_seq(row.number)
            created.append(e)
        db.flush()
        # Post through the engine; the retroactive negative-balance
        # check covers TK 1111 only (bank balances may not be entered
        # yet).
        js.post_batch(
            db,
            ctx,
            created,
            negative_accounts={CASH},
            action_label="nhập sổ quỹ",
        )
        for e in created:
            write_audit(
                db,
                ctx,
                action="entry.import",
                entity_type="entry",
                entity_id=e.id,
                summary=(
                    f"Nhập từ Excel {e.number} ngày {fmt_date(e.date)}: "
                    f"{format_vnd(e.total)}"
                ),
            )
        write_audit(
            db,
            ctx,
            action="import.confirm",
            entity_type="import",
            entity_id=token,
            summary=f"Nhập {len(created)} phiếu thu, chi từ file {filename}",
            after={
                "filename": filename,
                "imported": len(created),
                "error_rows": len(preview.error_rows),
                "skipped_examples": preview.skipped_examples,
                "total_in": preview.total_in,
                "total_out": preview.total_out,
                "opening_set": opening_set,
                "opening_balance": preview.opening_file,
                "closing_after": preview.closing_after,
            },
        )
    return preview


# ------------------------------------------------------ temporary files


def _paths(token: str) -> tuple[Path, Path]:
    if not TOKEN_RE.match(token or ""):
        raise NotFoundError("Không tìm thấy phiên nhập dữ liệu")
    base = get_settings().imports_dir
    return base / f"{token}.xlsx", base / f"{token}.json"


def save_upload(
    company_id: int, filename: str, data: bytes, user_id: int | None
) -> str:
    settings = get_settings()
    if not data:
        raise BusinessError("File rỗng")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise BusinessError(f"File vượt quá {settings.max_upload_mb} MB")
    if not (filename or "").lower().endswith(".xlsx") or not data.startswith(
        b"PK"
    ):
        raise BusinessError("Chỉ nhận file Excel .xlsx")
    settings.ensure_dirs()
    cleanup_old()
    token = secrets.token_hex(16)
    xlsx, meta = _paths(token)
    xlsx.write_bytes(data)
    meta.write_text(
        json.dumps(
            {
                "filename": Path(filename).name[:200],
                "user_id": user_id,
                "company_id": company_id,
                "uploaded_at": _now().isoformat(),
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return token


def load_upload(company_id: int, token: str) -> tuple[Path, str]:
    xlsx, meta = _paths(token)
    if not xlsx.exists() or not meta.exists():
        raise NotFoundError(
            "Phiên nhập dữ liệu đã hết hạn hoặc đã xử lý. "
            "Hãy tải file lên lại."
        )
    try:
        info = json.loads(meta.read_text(encoding="utf-8"))
    except (ValueError, OSError) as e:
        raise NotFoundError("Phiên nhập dữ liệu không hợp lệ") from e
    if info.get("company_id") != company_id:
        raise NotFoundError("Không tìm thấy phiên nhập dữ liệu")
    return xlsx, info.get("filename") or "import.xlsx"


def discard_upload(token: str, company_id: int | None = None) -> None:
    """Delete an import session.

    Pass ``company_id`` to delete only a session of that company.
    """
    xlsx, meta = _paths(token)
    if company_id is not None and meta.exists():
        try:
            owner = json.loads(meta.read_text(encoding="utf-8")).get(
                "company_id"
            )
        except (ValueError, OSError):
            owner = None
        if owner != company_id:
            raise NotFoundError("Không tìm thấy phiên nhập dữ liệu")
    for p in (xlsx, meta):
        p.unlink(missing_ok=True)


def cleanup_old() -> None:
    """Remove old temporary import files/dirs (of every module)."""
    from app.core.tempfiles import cleanup_imports

    cleanup_imports(TEMP_MAX_AGE / 3600)
