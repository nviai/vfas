"""Bank statements.

Reads CSV/XLSX files, maps columns, skips duplicates, auto-matches
lines with vouchers and creates draft vouchers from unmatched lines.

Sign convention: a statement credit is money INTO the account, which
the books record as Debit 112 (bank receipt, BC); a statement debit
is money OUT, recorded as Credit 112 (bank payment, BN).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import secrets
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app.core.audit import write_audit
from app.core.config import get_settings
from app.core.errors import BusinessError, NotFoundError
from app.core.money import MAX_MONEY, format_vnd
from app.features.bank.models import BankAccount
from app.features.bank.statement_models import (
    IGNORED,
    MATCHED,
    UNMATCHED,
    BankImportProfile,
    BankRule,
    BankStatement,
    BankStatementLine,
)
from app.features.journal import queries
from app.features.journal import service as js
from app.features.journal.models import POSTED, JournalEntry, JournalLine

FIELDS = {
    "date": "Ngày giao dịch",
    "description": "Diễn giải",
    "money_in": "Ghi có (tiền vào)",
    "money_out": "Ghi nợ (tiền ra)",
    "amount": "Số tiền (+ vào / − ra)",
    "balance": "Số dư",
    "reference": "Số tham chiếu",
}
# Keywords that identify column headers (diacritics removed,
# lowercase)
SYNONYMS = {
    "date": (
        "ngay giao dich",
        "ngay gd",
        "ngay hach toan",
        "ngay",
        "transaction date",
        "posting date",
        "date",
        "value date",
        "ngay hieu luc",
    ),
    "description": (
        "dien giai",
        "noi dung",
        "mo ta",
        "chi tiet giao dich",
        "description",
        "details",
        "narrative",
        "remark",
        "transaction details",
    ),
    "money_in": (
        "ghi co",
        "so tien ghi co",
        "phat sinh co",
        "tien vao",
        "credit",
        "credit amount",
        "co",
    ),
    "money_out": (
        "ghi no",
        "so tien ghi no",
        "phat sinh no",
        "tien ra",
        "debit",
        "debit amount",
        "no",
    ),
    "amount": ("so tien", "amount", "so tien giao dich"),
    "balance": ("so du", "balance", "so du cuoi", "running balance"),
    "reference": (
        "so tham chieu",
        "ma giao dich",
        "so but toan",
        "so ct",
        "reference",
        "ref",
        "transaction no",
        "ma gd",
        "so giao dich",
    ),
}
MAX_ROWS = 5000
MAX_COLS = 50
MATCH_DAYS = 3


def fold(text) -> str:
    s = (
        unicodedata.normalize("NFD", str(text or ""))
        .replace("đ", "d")
        .replace("Đ", "D")
    )
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


# -------------------------------------------------------- reading files


def read_rows(data: bytes, filename: str) -> list[list]:
    """Read a CSV/XLSX file into a list of rows (raw values)."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(
                io.BytesIO(data), data_only=True, read_only=True
            )
        except Exception as e:  # noqa: BLE001
            raise BusinessError("Không đọc được file Excel") from e
        try:
            ws = wb.worksheets[0]
            rows = []
            for values in ws.iter_rows(values_only=True, max_col=MAX_COLS):
                rows.append(list(values)[:MAX_COLS])
                if len(rows) > MAX_ROWS + 50:
                    break
            return rows
        finally:
            wb.close()
    if name.endswith(".xls"):
        raise BusinessError(
            "File .xls (Excel cũ) chưa hỗ trợ – mở bằng Excel và lưu lại "
            "dạng .xlsx hoặc .csv"
        )
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    try:
        for row in csv.reader(io.StringIO(text), dialect):
            rows.append(row[:MAX_COLS])
            if len(rows) >= MAX_ROWS + 50:
                break
    except csv.Error as e:
        raise BusinessError("File CSV không đúng định dạng") from e
    return rows


def guess_mapping(rows: list[list]) -> tuple[int, dict[str, int]]:
    """Find the header row and the column of each field.

    The header row is the one with the most cells matching a keyword.
    """
    best = (-1, 0, {})
    for i, row in enumerate(rows[:40]):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            h = fold(cell)
            if not h or len(h) > 60:
                continue
            for field_, keys in SYNONYMS.items():
                if field_ in mapping:
                    continue
                if any(h == k or (len(k) > 3 and k in h) for k in keys):
                    # "so tien ghi co" contains "so tien": prefer the
                    # credit/debit columns over "amount"
                    if field_ == "amount" and any(
                        k in h for k in ("ghi co", "ghi no", "credit", "debit")
                    ):
                        continue
                    mapping[field_] = col
                    break
        score = len(mapping)
        if score > best[1]:
            best = (i, score, mapping)
    return best[0], best[2]


def parse_amount(v) -> int | None:
    """Parse an amount into an int (VND); return None if empty.

    Accepts '1,000,000', '1.000.000', '1000000.00', '-50,000' or a
    number.
    """
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v if abs(v) <= MAX_MONEY else None
    if isinstance(v, float):
        return (
            int(round(v)) if math.isfinite(v) and abs(v) <= MAX_MONEY else None
        )
    if isinstance(v, Decimal):
        return (
            int(v.to_integral_value())
            if v.is_finite() and abs(v) <= MAX_MONEY
            else None
        )
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if len(s) > 40:
        return None
    s = re.sub(r"(?i)(vnd|vnđ|đ|d)$", "", s)
    if not s or s in ("-", "--"):
        return None
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("-()+")
    if not re.fullmatch(r"[\d.,]+", s):
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        thou = "." if dec == "," else ","
        s = s.replace(thou, "").replace(dec, ".")
    elif s.count(",") > 1 or s.count(".") > 1:
        s = s.replace(",", "").replace(".", "")
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        head, tail = s.split(sep)
        s = head + tail if len(tail) == 3 else head + "." + tail
    try:
        n = Decimal(s)
    except InvalidOperation:
        return None
    if not n.is_finite() or abs(n) > MAX_MONEY:
        return None
    n = int(n.quantize(Decimal(1)))
    return -n if neg else n


def parse_day(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip().split(" ")[0].split("T")[0]
    for fmt in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d/%m/%y",
        "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


@dataclass
class ParsedLine:
    row_no: int
    date: date | None
    description: str
    reference: str
    money_in: int
    money_out: int
    balance: int | None
    error: str = ""


@dataclass
class ParseResult:
    header_row: int
    mapping: dict[str, int]
    lines: list[ParsedLine] = field(default_factory=list)
    skipped: int = 0

    @property
    def valid(self) -> list[ParsedLine]:
        return [ln for ln in self.lines if not ln.error]

    @property
    def errors(self) -> list[ParsedLine]:
        return [ln for ln in self.lines if ln.error]


def apply_mapping(
    rows: list[list], header_row: int, mapping: dict[str, int]
) -> ParseResult:
    res = ParseResult(header_row=header_row, mapping=mapping)
    if "date" not in mapping or not (
        ("money_in" in mapping and "money_out" in mapping)
        or "amount" in mapping
    ):
        return res

    def cell(row, key):
        idx = mapping.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    for i, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        if not any(c not in (None, "") for c in row):
            continue
        raw_date = cell(row, "date")
        d = parse_day(raw_date)
        desc = str(cell(row, "description") or "").strip()
        if d is None:
            # Total rows, footers, etc. are skipped silently
            res.skipped += 1
            continue
        if "amount" in mapping and "money_in" not in mapping:
            amt = parse_amount(cell(row, "amount"))
            m_in, m_out = (amt, 0) if (amt or 0) > 0 else (0, -(amt or 0))
        else:
            m_in = abs(parse_amount(cell(row, "money_in")) or 0)
            m_out = abs(parse_amount(cell(row, "money_out")) or 0)
        ref = str(cell(row, "reference") or "").strip()[:100]
        bal = (
            parse_amount(cell(row, "balance"))
            if "balance" in mapping
            else None
        )
        line = ParsedLine(
            row_no=i,
            date=d,
            description=desc[:2000],
            reference=ref,
            money_in=m_in,
            money_out=m_out,
            balance=bal,
        )
        if m_in and m_out:
            line.error = "Vừa có tiền vào vừa có tiền ra"
        elif not m_in and not m_out:
            line.error = "Không có số tiền"
        res.lines.append(line)
        if len(res.lines) >= MAX_ROWS:
            break
    return res


def fingerprint(bank_account_id: int, ln: ParsedLine) -> str:
    raw = (
        f"{bank_account_id}|{ln.date}|{ln.money_in}|{ln.money_out}"
        f"|{fold(ln.description)}|{ln.reference}|{ln.balance}"
    )
    return hashlib.sha256(raw.encode()).hexdigest()


# ------------------------------------------------------ temporary files

TOKEN_RE = re.compile(r"^[0-9a-f]{32}$")


def _paths(token: str) -> tuple[Path, Path]:
    if not TOKEN_RE.match(token or ""):
        raise NotFoundError("Không tìm thấy phiên nhập sao kê")
    base = get_settings().imports_dir
    return base / f"bank-{token}.bin", base / f"bank-{token}.json"


def save_upload(
    company_id: int, bank_account_id: int, filename: str, data: bytes
) -> str:
    settings = get_settings()
    if not data:
        raise BusinessError("File rỗng")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise BusinessError(f"File vượt quá {settings.max_upload_mb} MB")
    if (
        not (filename or "")
        .lower()
        .endswith((".csv", ".xlsx", ".txt", ".xls"))
    ):
        raise BusinessError("Chỉ nhận file sao kê .xlsx hoặc .csv")
    settings.ensure_dirs()
    from app.core.tempfiles import cleanup_imports_throttled

    cleanup_imports_throttled()
    token = secrets.token_hex(16)
    raw, meta = _paths(token)
    raw.write_bytes(data)
    meta.write_text(
        json.dumps(
            {
                "company_id": company_id,
                "bank_account_id": bank_account_id,
                "filename": Path(filename).name[:200],
            }
        ),
        encoding="utf-8",
    )
    return token


def load_upload(company_id: int, token: str) -> tuple[bytes, dict]:
    raw, meta = _paths(token)
    if not raw.exists() or not meta.exists():
        raise NotFoundError("Phiên nhập sao kê đã hết hạn hoặc đã xử lý")
    info = json.loads(meta.read_text(encoding="utf-8"))
    if info.get("company_id") != company_id:
        raise NotFoundError("Không tìm thấy phiên nhập sao kê")
    return raw.read_bytes(), info


def discard_upload(token: str, company_id: int | None = None) -> None:
    """Delete a statement import session.

    Pass ``company_id`` to delete only a session of that company.
    """
    raw, meta = _paths(token)
    if company_id is not None and meta.exists():
        try:
            owner = json.loads(meta.read_text(encoding="utf-8")).get(
                "company_id"
            )
        except (ValueError, OSError):
            owner = None
        if owner != company_id:
            raise NotFoundError("Không tìm thấy phiên nhập sao kê")
    for p in (raw, meta):
        p.unlink(missing_ok=True)


# ---------------------------------------------------- saving statements


def get_profile(
    db: Session, company_id: int, bank_account_id: int
) -> dict | None:
    p = db.scalar(
        select(BankImportProfile).where(
            BankImportProfile.company_id == company_id,
            BankImportProfile.bank_account_id == bank_account_id,
        )
    )
    return dict(p.mapping) if p and p.mapping else None


def save_profile(
    db: Session, company_id: int, bank_account_id: int, mapping: dict
) -> None:
    p = db.scalar(
        select(BankImportProfile).where(
            BankImportProfile.company_id == company_id,
            BankImportProfile.bank_account_id == bank_account_id,
        )
    )
    if p is None:
        p = BankImportProfile(
            company_id=company_id, bank_account_id=bank_account_id
        )
        db.add(p)
    p.mapping = mapping
    p.updated_at = datetime.now(timezone.utc)


def import_statement(
    db: Session,
    ctx,
    bank: BankAccount,
    filename: str,
    result: ParseResult,
    header_row: int,
    mapping: dict[str, int],
) -> tuple[BankStatement, int, int]:
    """Save the valid lines, skipping lines imported before.

    Return (statement, number of new lines, number of duplicates).
    """
    if bank.company_id != ctx.company_id:
        raise NotFoundError()
    valid = result.valid
    if not valid:
        raise BusinessError("Không có dòng giao dịch hợp lệ trong file")
    existing = set(
        db.scalars(
            select(BankStatementLine.fingerprint).where(
                BankStatementLine.company_id == ctx.company_id,
                BankStatementLine.bank_account_id == bank.id,
            )
        )
    )
    st = BankStatement(
        company_id=ctx.company_id,
        bank_account_id=bank.id,
        filename=filename[:255],
        imported_by=ctx.user_id,
        imported_at=datetime.now(timezone.utc),
    )
    db.add(st)
    db.flush()
    new = dup = 0
    seen: set[str] = set()
    for ln in sorted(valid, key=lambda x: (x.date, x.row_no)):
        fp = fingerprint(bank.id, ln)
        if fp in existing or fp in seen:
            dup += 1
            continue
        seen.add(fp)
        new += 1
        st.lines.append(
            BankStatementLine(
                company_id=ctx.company_id,
                bank_account_id=bank.id,
                line_no=new,
                date=ln.date,
                description=ln.description,
                reference=ln.reference,
                money_in=ln.money_in,
                money_out=ln.money_out,
                balance=ln.balance,
                fingerprint=fp,
                status=UNMATCHED,
            )
        )
    if new:
        dates = [ln.date for ln in valid]
        st.period_start, st.period_end = min(dates), max(dates)
        with_bal = [
            ln
            for ln in sorted(valid, key=lambda x: (x.date, x.row_no))
            if ln.balance is not None
        ]
        st.closing_balance = with_bal[-1].balance if with_bal else None
    st.line_count = new
    save_profile(
        db,
        ctx.company_id,
        bank.id,
        {"header_row": header_row, **{k: v for k, v in mapping.items()}},
    )
    write_audit(
        db,
        ctx,
        action="bank.statement_import",
        entity_type="bank_statement",
        entity_id=st.id,
        summary=(
            f"Nhập sao kê {bank.label}: {new} giao dịch mới, "
            f"{dup} trùng ({filename})"
        ),
    )
    db.flush()
    if not new:
        db.delete(st)
    db.commit()
    if new:
        auto_match(db, ctx, bank.id)
    return st, new, dup


# ------------------------------------------------------------- matching


def _bank_side(bank: BankAccount, incoming: bool):
    jl = JournalLine
    if incoming:
        return and_(
            jl.debit_account == bank.account_code,
            jl.debit_bank_account_id == bank.id,
        )
    return and_(
        jl.credit_account == bank.account_code,
        jl.credit_bank_account_id == bank.id,
    )


def _taken_line_ids(db: Session, company_id: int) -> set[int]:
    return set(
        db.scalars(
            select(BankStatementLine.matched_line_id).where(
                BankStatementLine.company_id == company_id,
                BankStatementLine.matched_line_id.is_not(None),
            )
        )
    )


def candidates(
    db: Session, company_id: int, sl: BankStatementLine, *, days: int = 30
) -> list[tuple]:
    """Return posted voucher lines that could match a statement line.

    Candidates use the same bank account, direction and amount, dated
    within +/- ``days`` days.
    """
    bank = db.get(BankAccount, sl.bank_account_id)
    taken = _taken_line_ids(db, company_id) - {sl.matched_line_id}
    jl, je = JournalLine, JournalEntry
    rows = db.execute(
        select(jl, je)
        .join(je, je.id == jl.entry_id)
        .where(
            je.company_id == company_id,
            je.status == POSTED,
            _bank_side(bank, sl.incoming),
            jl.amount == sl.amount,
            je.date >= sl.date - timedelta(days=days),
            je.date <= sl.date + timedelta(days=days),
        )
        .order_by(je.date, jl.id)
    ).all()
    return [(ln, e) for ln, e in rows if ln.id not in taken]


def _score(sl: BankStatementLine, ln: JournalLine, e: JournalEntry) -> tuple:
    text = fold(
        " ".join(
            [
                e.ref_no or "",
                e.number or "",
                e.description or "",
                ln.description or "",
            ]
        )
    )
    ref_hit = bool(sl.reference) and fold(sl.reference) in text
    num_hit = bool(e.number) and fold(e.number) in fold(sl.description)
    return (0 if (ref_hit or num_hit) else 1, abs((e.date - sl.date).days))


def auto_match(db: Session, ctx, bank_account_id: int) -> int:
    """Auto-match the unmatched statement lines of a bank account.

    A match needs the same amount and a date gap of at most 3 days;
    candidates with a matching reference number are preferred.
    """
    lines = list(
        db.scalars(
            select(BankStatementLine)
            .where(
                BankStatementLine.company_id == ctx.company_id,
                BankStatementLine.bank_account_id == bank_account_id,
                BankStatementLine.status == UNMATCHED,
            )
            .order_by(BankStatementLine.date, BankStatementLine.id)
        )
    )
    matched = 0
    for sl in lines:
        cands = [
            (ln, e)
            for ln, e in candidates(db, ctx.company_id, sl, days=MATCH_DAYS)
        ]
        if not cands:
            continue
        ln, e = min(cands, key=lambda c: _score(sl, c[0], c[1]))
        _set_match(sl, ln, e, ctx.user_id)
        db.flush()
        matched += 1
    if matched:
        write_audit(
            db,
            ctx,
            action="bank.auto_match",
            entity_type="bank_account",
            entity_id=bank_account_id,
            summary=f"Tự khớp {matched} giao dịch sao kê",
        )
    db.commit()
    return matched


def _set_match(
    sl: BankStatementLine,
    ln: JournalLine | None,
    e: JournalEntry | None,
    user_id,
) -> None:
    sl.matched_line_id = ln.id if ln else None
    sl.matched_entry_id = e.id if e else None
    sl.status = MATCHED if ln else UNMATCHED
    sl.matched_by = user_id if ln else None
    sl.matched_at = datetime.now(timezone.utc) if ln else None


def get_line(db: Session, company_id: int, line_id: int) -> BankStatementLine:
    sl = db.get(BankStatementLine, line_id)
    if sl is None or sl.company_id != company_id:
        raise NotFoundError("Không tìm thấy dòng sao kê")
    return sl


def match(
    db: Session, ctx, line_id: int, journal_line_id: int
) -> BankStatementLine:
    sl = get_line(db, ctx.company_id, line_id)
    ok = {
        ln.id: (ln, e)
        for ln, e in candidates(db, ctx.company_id, sl, days=366)
    }
    if journal_line_id not in ok:
        raise BusinessError(
            "Chứng từ không khớp số tiền / tài khoản ngân hàng / chiều tiền "
            "với dòng sao kê"
        )
    ln, e = ok[journal_line_id]
    _set_match(sl, ln, e, ctx.user_id)
    write_audit(
        db,
        ctx,
        action="bank.match",
        entity_type="bank_statement_line",
        entity_id=sl.id,
        summary=(
            f"Khớp giao dịch {sl.date:%d/%m/%Y} {format_vnd(sl.amount)} "
            f"với {e.number}"
        ),
    )
    db.commit()
    return sl


def unmatch(db: Session, ctx, line_id: int) -> BankStatementLine:
    sl = get_line(db, ctx.company_id, line_id)
    _set_match(sl, None, None, ctx.user_id)
    sl.status = UNMATCHED
    write_audit(
        db,
        ctx,
        action="bank.unmatch",
        entity_type="bank_statement_line",
        entity_id=sl.id,
        summary=(
            f"Bỏ khớp giao dịch {sl.date:%d/%m/%Y} {format_vnd(sl.amount)}"
        ),
    )
    db.commit()
    return sl


def ignore(
    db: Session, ctx, line_id: int, note: str = ""
) -> BankStatementLine:
    sl = get_line(db, ctx.company_id, line_id)
    if sl.status == MATCHED:
        raise BusinessError("Dòng đã khớp – bỏ khớp trước")
    sl.status = IGNORED if sl.status != IGNORED else UNMATCHED
    sl.note = (note or "").strip()[:300]
    write_audit(
        db,
        ctx,
        action="bank.ignore",
        entity_type="bank_statement_line",
        entity_id=sl.id,
        summary=("Bỏ qua" if sl.status == IGNORED else "Theo dõi lại")
        + f" giao dịch {sl.date:%d/%m/%Y} {format_vnd(sl.amount)}",
    )
    db.commit()
    return sl


# ------------------------------------- suggestions and voucher creation


def rules(db: Session, company_id: int) -> list[BankRule]:
    return list(
        db.scalars(
            select(BankRule)
            .where(BankRule.company_id == company_id)
            .order_by(BankRule.priority, BankRule.id)
        )
    )


def suggest(
    db: Session, company_id: int, sl: BankStatementLine
) -> BankRule | None:
    text = fold(sl.description) + " " + fold(sl.reference)
    direction = "in" if sl.incoming else "out"
    for r in rules(db, company_id):
        if not r.is_active or r.direction not in ("any", direction):
            continue
        if fold(r.keyword) and fold(r.keyword) in text:
            return r
    return None


def create_draft(db: Session, ctx, line_id: int) -> JournalEntry:
    """Create a DRAFT bank receipt/payment from an unmatched line.

    The offsetting account is suggested by the keyword rules.
    """
    sl = get_line(db, ctx.company_id, line_id)
    if sl.status == MATCHED:
        raise BusinessError("Dòng sao kê đã khớp với chứng từ")
    existing = db.scalar(
        select(JournalEntry).where(
            JournalEntry.company_id == ctx.company_id,
            JournalEntry.status != "cancelled",
            JournalEntry.meta["statement_line_id"].as_integer() == sl.id,
        )
    )
    if existing is not None:
        return existing
    bank = db.get(BankAccount, sl.bank_account_id)
    rule = suggest(db, ctx.company_id, sl)
    contra = rule.account if rule else ""
    partner_id = rule.partner_id if rule else None
    desc = (
        rule.description
        if rule and rule.description
        else sl.description or "Giao dịch ngân hàng"
    )[:500]
    if sl.incoming:
        line = js.LineIn(
            debit_account=bank.account_code,
            credit_account=contra,
            amount=sl.amount,
            description=desc,
            debit_bank_account_id=bank.id,
            credit_partner_id=partner_id,
        )
    else:
        line = js.LineIn(
            debit_account=contra,
            credit_account=bank.account_code,
            amount=sl.amount,
            description=desc,
            credit_bank_account_id=bank.id,
            debit_partner_id=partner_id,
        )
    inp = js.EntryIn(
        doc_type="BC" if sl.incoming else "BN",
        date=sl.date,
        doc_date=sl.date,
        lines=[line],
        description=desc,
        partner_id=partner_id,
        ref_no=sl.reference,
        source="bank",
        meta={"form": "money", "statement_line_id": sl.id},
    )
    return js.create_entry(db, ctx, inp, allow_incomplete=True)


def match_posted_entry(db: Session, ctx, e: JournalEntry) -> None:
    """Hook: match a posted voucher with its source statement line."""
    line_id = (e.meta or {}).get("statement_line_id")
    if not line_id:
        return
    sl = db.get(BankStatementLine, int(line_id))
    if sl is None or sl.company_id != e.company_id or sl.status == MATCHED:
        return
    bank = db.get(BankAccount, sl.bank_account_id)
    for ln in e.lines:
        side_ok = (
            (
                ln.debit_account == bank.account_code
                and ln.debit_bank_account_id == bank.id
            )
            if sl.incoming
            else (
                ln.credit_account == bank.account_code
                and ln.credit_bank_account_id == bank.id
            )
        )
        if side_ok and ln.amount == sl.amount:
            _set_match(sl, ln, e, ctx.user_id)
            return


def release_cancelled_entry(db: Session, ctx, e: JournalEntry) -> None:
    """Hook: a cancelled matched voucher unmatches its lines."""
    for sl in db.scalars(
        select(BankStatementLine).where(
            BankStatementLine.company_id == e.company_id,
            BankStatementLine.matched_entry_id == e.id,
        )
    ):
        _set_match(sl, None, None, None)


# ------------------------------------------------------- reconciliation


@dataclass
class Reconciliation:
    bank: BankAccount
    as_of: date | None
    statement_balance: int | None
    book_balance: int
    unmatched_lines: list[BankStatementLine]
    unmatched_book: list[tuple[JournalLine, JournalEntry]]
    matched_count: int
    total_count: int

    @property
    def difference(self) -> int | None:
        return (
            None
            if self.statement_balance is None
            else self.statement_balance - self.book_balance
        )


def reconciliation(
    db: Session, company_id: int, bank: BankAccount
) -> Reconciliation:
    lines = list(
        db.scalars(
            select(BankStatementLine)
            .where(
                BankStatementLine.company_id == company_id,
                BankStatementLine.bank_account_id == bank.id,
            )
            .order_by(BankStatementLine.date, BankStatementLine.id)
        )
    )
    with_bal = [sl for sl in lines if sl.balance is not None]
    last = with_bal[-1] if with_bal else None
    as_of = last.date if last else (lines[-1].date if lines else None)
    book = (
        queries.balance(
            db,
            company_id,
            bank.account_code,
            upto=as_of,
            exact=True,
            bank_account_id=bank.id,
        )
        if as_of
        else queries.balance(
            db,
            company_id,
            bank.account_code,
            exact=True,
            bank_account_id=bank.id,
        )
    )
    unmatched = [sl for sl in lines if sl.status == UNMATCHED]
    unmatched_book: list = []
    if lines:
        start = lines[0].date
        taken = _taken_line_ids(db, company_id)
        jl, je = JournalLine, JournalEntry
        rows = db.execute(
            select(jl, je)
            .join(je, je.id == jl.entry_id)
            .where(
                je.company_id == company_id,
                je.status == POSTED,
                je.date >= start,
                je.date <= (as_of or start),
                or_(_bank_side(bank, True), _bank_side(bank, False)),
            )
            .order_by(je.date, je.posting_seq)
        ).all()
        unmatched_book = [(ln, e) for ln, e in rows if ln.id not in taken]
    return Reconciliation(
        bank=bank,
        as_of=as_of,
        statement_balance=last.balance if last else None,
        book_balance=book,
        unmatched_lines=unmatched,
        unmatched_book=unmatched_book,
        matched_count=sum(1 for sl in lines if sl.status == MATCHED),
        total_count=len(lines),
    )


def list_lines(
    db: Session,
    company_id: int,
    *,
    bank_account_id: int | None = None,
    status: str = "",
    statement_id: int | None = None,
    limit: int = 500,
) -> list[BankStatementLine]:
    stmt = select(BankStatementLine).where(
        BankStatementLine.company_id == company_id
    )
    if bank_account_id:
        stmt = stmt.where(BankStatementLine.bank_account_id == bank_account_id)
    if status in (UNMATCHED, MATCHED, IGNORED):
        stmt = stmt.where(BankStatementLine.status == status)
    if statement_id:
        stmt = stmt.where(BankStatementLine.statement_id == statement_id)
    return list(
        db.scalars(
            stmt.order_by(
                BankStatementLine.date.desc(), BankStatementLine.id.desc()
            ).limit(limit)
        )
    )


def save_rule(
    db: Session,
    ctx,
    *,
    keyword: str,
    account: str,
    direction: str = "any",
    partner_id: int | None = None,
    description: str = "",
    priority: int = 100,
) -> BankRule:
    from app.features.accounts.models import Account

    keyword = (keyword or "").strip()[:100]
    if not keyword:
        raise BusinessError("Nhập từ khóa")
    acc = db.scalar(
        select(Account).where(
            Account.company_id == ctx.company_id,
            Account.code == (account or "").strip(),
        )
    )
    if acc is None or not acc.is_postable:
        raise BusinessError("TK đối ứng phải là tài khoản chi tiết")
    if direction not in ("in", "out", "any"):
        direction = "any"
    r = BankRule(
        company_id=ctx.company_id,
        keyword=keyword,
        direction=direction,
        account=acc.code,
        partner_id=partner_id or None,
        description=(description or "").strip()[:300],
        priority=priority,
        is_active=True,
    )
    db.add(r)
    db.flush()
    write_audit(
        db,
        ctx,
        action="bank.rule_create",
        entity_type="bank_rule",
        entity_id=r.id,
        summary=f"Thêm luật sao kê '{keyword}' → {acc.code}",
    )
    db.commit()
    return r


def delete_rule(db: Session, ctx, rule_id: int) -> None:
    r = db.get(BankRule, rule_id)
    if r is None or r.company_id != ctx.company_id:
        raise NotFoundError()
    write_audit(
        db,
        ctx,
        action="bank.rule_delete",
        entity_type="bank_rule",
        entity_id=r.id,
        summary=f"Xóa luật sao kê '{r.keyword}'",
    )
    db.delete(r)
    db.commit()
